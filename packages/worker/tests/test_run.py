"""Сквозные проверки сборщика на примерах из examples/: здесь модули работают вместе."""

import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from autogenerator.contracts import ErrorCode, Period, PreviewRequest, RunRequest, SourceSpec
from autogenerator.contracts.yaml_io import load_model_list, load_yaml
from autogenerator.worker import load_scenario, preview, run

EXAMPLE = Path(__file__).resolve().parents[3] / "examples" / "sales"


def request(tmp_path: Path, **kw) -> RunRequest:
    data = EXAMPLE / "data"
    params = {
        "scenario": load_scenario(load_yaml(EXAMPLE / "scenario.yaml")),
        "sources": load_model_list(SourceSpec, EXAMPLE / "sources.yaml"),
        "inputs": {
            "sales": [str(p) for p in sorted((data / "sales").glob("*.csv"))],
            "plan": [str(p) for p in sorted((data / "plan").glob("*.xlsx"))],
        },
        "theme": str(EXAMPLE.parent / "templates" / "synthetic.pptx"),
        "output_dir": str(tmp_path),
    }
    params.update(kw)
    return RunRequest(**params)


def slide_texts(path: str) -> list[str]:
    prs = Presentation(path)
    return [" | ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


def test_example_report(tmp_path: Path):
    res = run(request(tmp_path))
    assert res.ok, res.errors
    assert res.period.key == "2026-03"
    assert res.output_path == str(tmp_path / "Отчёт_продажи_2026-03.pptx")
    assert res.slides == 7
    texts = slide_texts(res.output_path)
    assert texts[0].startswith("Продажи: март 2026")
    assert "План месяца выполнен на 103,2%" in texts[1]
    assert "К прошлому месяцу: +8,3%" in texts[1]
    assert texts[6].startswith("План-факт: март 2026")
    # Три месяца продаж с переименованными и переставленными столбцами прочитаны как один источник.
    assert [u["period"] for u in res.inputs["sales"]] == ["2026-01", "2026-02", "2026-03"]
    assert res.inputs["plan"][0]["period"] == "2026-Q1"
    assert res.workdir is None


def test_other_period_and_workdir(tmp_path: Path):
    work = tmp_path / "work"
    res = run(request(tmp_path, period=Period.parse("2026-02"), workdir=str(work)))
    assert res.ok, res.errors
    assert slide_texts(res.output_path)[0].startswith("Продажи: февраль 2026")
    # Промежуточные данные остаются: каждый модуль можно перезапустить на них отдельно.
    for p in [
        "theme/manifest.json",
        "manifests/sales.json",
        "engine/inputs/sales.parquet",
        "engine/outputs/metrics.json",
        "run.json",
    ]:
        assert (work / p).exists(), p


def test_missing_column_blocks_the_run(tmp_path: Path):
    bad = tmp_path / "bad.csv"
    bad.write_text("Дата заказа;Регион\n01.04.2026;Москва\n", encoding="utf-8")
    req = request(tmp_path)
    req.inputs["sales"] = [bad]
    res = run(req)
    assert not res.ok and res.output_path is None
    assert res.errors[0].code == ErrorCode.SCHEMA_BLOCKED
    assert "Сумма" in res.errors[0].message


def test_cast_errors_need_confirmation(tmp_path: Path):
    src = EXAMPLE / "data" / "sales" / "Продажи_2026-03.csv"
    lines = src.read_bytes().decode("cp1251").splitlines()
    for i in range(1, 20):
        lines[i] = lines[i].rsplit(";", 1)[0] + ";н/д"
    broken = tmp_path / "Продажи_2026-03.csv"
    broken.write_bytes("\r\n".join(lines).encode("cp1251"))
    req = request(tmp_path)
    req.inputs["sales"] = [*req.inputs["sales"][:2], str(broken)]
    res = run(req)
    assert res.errors and res.errors[0].code == ErrorCode.CAST_REVIEW
    assert "--accept-cast-errors" in res.errors[0].message

    res = run(req.model_copy(update={"accept_cast_errors": True}))
    assert res.ok
    assert any("принято с ошибками" in w.message for w in res.warnings)


def test_broken_scenario_is_reported_before_reading(tmp_path: Path):
    req = request(tmp_path)
    req.scenario.metrics[0].column = "nope"
    req.inputs["sales"] = ["/nonexistent.csv"]
    res = run(req)
    assert not res.ok
    assert any("nope" in e.message for e in res.errors)


@pytest.fixture(autouse=True)
def _no_leftovers(tmp_path: Path):
    yield
    shutil.rmtree(tmp_path / "work", ignore_errors=True)


def preview_request(target: str, **kw) -> PreviewRequest:
    r = request(Path("."))
    return PreviewRequest(scenario=r.scenario, sources=r.sources, inputs=r.inputs, target=target, **kw)


def test_preview_nodes_and_cache(tmp_path: Path):
    cache = tmp_path / "cache"
    res = preview(preview_request("input:sales/step:positive_only", cache_dir=str(cache), temp_dir=str(tmp_path)))
    assert not res.errors, res.errors
    assert res.period.key == "2026-03"
    steps = {s.id: s for s in res.steps}
    assert steps["dedupe_orders"].rows_after == steps["positive_only"].rows_before
    assert list(steps) == ["dedupe_orders", "positive_only"]  # шаги после выбранного не выполняются
    assert res.total_rows == steps["positive_only"].rows_after and len(res.rows) == 20
    # Повторное превью берёт готовый результат из кэша: те же числа.
    again = preview(preview_request("input:sales/step:positive_only", cache_dir=str(cache), temp_dir=str(tmp_path)))
    assert again.total_rows == res.total_rows and again.rows == res.rows
    assert any(cache.rglob("*.parquet"))

    res = preview(preview_request("dataset:by_region", rows=2))
    assert res.total_rows == 5 and len(res.rows) == 2
    assert {"revenue_prev_change_pct", "share"} <= {c.name for c in res.columns}
    res = preview(preview_request("metric:plan_done"))
    assert res.value == pytest.approx(1.032, abs=1e-3)
