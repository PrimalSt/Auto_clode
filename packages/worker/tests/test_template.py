"""Приёмочный тест шаблона (ARCHITECTURE.md, раздел 13).

Тест сам привязывает все метки, графики и таблицы слайдов-образцов к пробным значениям
(данные и сценарий пользователя ему не нужны), собирает отчёт и проверяет, что шаблон
заполнен без потерь:

- все метки заменены с тем же оформлением, кроме языка проверки правописания (он
  выставляется по значению), и «{{» не осталось;
- картинки, шрифты, мастера, макеты, темы и незаполняемые фигуры слайдов не изменились;
- у графиков то же число серий по группам, нет осей без графиков, а подписи линии
  комбинированных графиков не ложатся на подписи столбцов и друг на друга — на нескольких
  случайных наборах данных (значения шаблона, умноженные на случайные множители);
- пропущенный слайд удалён из файла;
- без одной привязки сборка останавливается со списком слайдов и фигур;
- файл открывается в PowerPoint, если он есть на компьютере.

По умолчанию — синтетический шаблон из examples/templates. Свой шаблон (сам он в репозиторий
не кладётся): ``agen test worker --template шаблон.pptx``.
"""

from __future__ import annotations

import copy
import hashlib
import os
import posixpath
import random
import shutil
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pyarrow as pa
import pytest
from lxml import etree
from pptx import Presentation
from pptx.opc.package import XmlPart

from autogenerator.blocks_std.chart_xml import check_chart, group_series, groups
from autogenerator.blocks_std.label_bands import label_overlaps
from autogenerator.contracts import (
    ChartInfo,
    EngineResult,
    IssueLevel,
    MarkerInfo,
    Period,
    RenderResult,
    ScenarioSpec,
    TemplateSlideInfo,
    ThemeManifest,
)
from autogenerator.contracts.ooxml import A, P, R, all_text, iter_markers, iter_shapes
from autogenerator.plugin_host import PluginRegistry
from autogenerator.render import build_presentation
from autogenerator.theme import import_template

SYNTHETIC = Path(__file__).resolve().parents[3] / "examples" / "templates" / "synthetic.pptx"
TEMPLATE = Path(os.environ.get("AGEN_TEST_TEMPLATE") or SYNTHETIC)
PERIOD = Period.parse("2026-03")
SEEDS = (1, 2, 3)
# Язык проверки правописания выставляется по значению, флаги ошибки и правки снимаются.
FREE_RUN_ATTRS = {"lang", "altLang", "err", "dirty"}
NOT_SLIDE_OWNED = ("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/", "ppt/notesMasters/", "ppt/handoutMasters/")


@dataclass
class Plan:
    """Сценарий, привязанный к пробным значениям, и что ожидать в отчёте."""

    scenario: ScenarioSpec
    engine: EngineResult
    slides: list[TemplateSlideInfo]
    values: dict[str, tuple[TemplateSlideInfo, MarkerInfo]] = field(default_factory=dict)
    skipped: TemplateSlideInfo | None = None
    categories: dict[tuple[int, int], int] = field(default_factory=dict)


@dataclass
class Built:
    plan: Plan
    result: RenderResult
    path: Path


# --- шаблон и пробные данные --------------------------------------------------------------


def template_slides(prs: Any) -> dict[int, Any]:
    lst = prs.slides._sldIdLst
    return {int(e.get("id")): prs.part.related_part(e.get(f"{R}id")).slide for e in lst}


def shapes_by_id(slide: Any) -> dict[int, Any]:
    out: dict[int, Any] = {}

    def walk(shapes: Any) -> None:
        for sh in shapes:
            out[sh.shape_id] = sh
            if sh.shape_type == 6:  # группа
                walk(sh.shapes)

    walk(slide.shapes)
    return out


def jitter(rnd: random.Random, v: float | None) -> float | None:
    return None if v is None else round(v * rnd.uniform(0.5, 1.5), 6)


def chart_data(
    rnd: random.Random, chart: Any, info: ChartInfo
) -> tuple[list[str], list[tuple[str, list[float | None]]]]:
    """Категории и серии графика шаблона в порядке python-pptx (так их ждёт ``chart_fill``) со
    случайными множителями; у графика без данных — случайные значения. Демо-данные шаблона
    (4,3; 2,5; …) у серий с подписями в процентах переводятся в доли, у 100%-гистограмм доли
    категории в сумме дают единицу — как у настоящих данных."""
    plots = list(chart.plots)
    cats = [str(c) for c in plots[0].categories] if plots else []
    series = [(s.name or "", list(s.values)) for p in plots for s in p.series]
    n = max([len(cats), *(len(v) for _, v in series)])
    if n == 0:
        n = 4
    cats = (cats + [f"К{i + 1}" for i in range(len(cats), n)])[:n]
    slots = [(g, si) for g in info.groups for si in g.series]
    if len(slots) != len(series):
        slots = []
    out: list[tuple[str, list[float | None]]] = []
    for k, (name, values) in enumerate(series):
        values = (values + [None] * n)[:n]
        if all(v is None or v == 0 for v in values):
            values = [rnd.uniform(1, 100) for _ in range(n)]
        vals = [jitter(rnd, v) for v in values]
        if slots and "%" in (slots[k][1].number_format or "") + (slots[k][0].label_format or ""):
            top = max((abs(v) for v in vals if v is not None), default=0.0)
            if top > 1:
                vals = [None if v is None else round(v / top * rnd.uniform(0.3, 0.95), 4) for v in vals]
        out.append((name, vals))
    for g in info.groups if slots else []:
        fmts = (g.label_format or "") + "".join(si.number_format or "" for si in g.series)
        if g.grouping != "percentStacked" or "%" not in fmts:
            continue
        idx = [k for k, (gg, _) in enumerate(slots) if gg is g]
        for i in range(n):
            total = sum(abs(out[k][1][i] or 0.0) for k in idx)
            for k in idx:
                v = out[k][1][i]
                out[k][1][i] = None if v is None or not total else round(v / total, 4)
    return cats, out


def make_plan(theme: ThemeManifest, seed: int) -> Plan:
    rnd = random.Random(seed)
    live = template_slides(Presentation(theme.pptx_path))
    usable = [s for s in theme.slides if all(m.replaceable for m in s.markers)]
    # Последний слайд-образец пропускается: он должен исчезнуть из файла.
    skipped = usable[-1] if len(usable) > 1 else None
    slides = [s for s in usable if s is not skipped]
    specs: list[dict[str, Any]] = []
    datasets: dict[str, pa.Table] = {}
    plan_values: dict[str, tuple[TemplateSlideInfo, MarkerInfo]] = {}
    categories: dict[tuple[int, int], int] = {}
    for s in slides:
        markers: dict[str, Any] = {}
        for m in s.markers:
            value = f"V{len(plan_values) + 1:04d}V"
            markers[m.key] = {"value": value}
            plan_values[value] = (s, m)
        shapes = shapes_by_id(live[s.slide_id])
        blocks: list[dict[str, Any]] = []
        for c in s.charts:
            name = f"chart_{s.slide_id}_{c.shape_id}"
            cats, series = chart_data(rnd, shapes[c.shape_id].chart, c)
            columns: dict[str, list[Any]] = {"cat": cats}
            spec_series: list[Any] = []
            for j, (sname, values) in enumerate(series):
                columns[f"s{j}"] = values
                spec_series.append({"column": f"s{j}", "name": sname} if sname else f"s{j}")
            datasets[name] = pa.table(columns)
            categories[(s.slide_id, c.shape_id)] = len(cats)
            blocks.append(
                {"type": "chart_fill", "shape": c.shape_id, "dataset": name, "categories": "cat", "series": spec_series}
            )
        marked = {m.shape_id for m in s.markers}
        keep = []
        for t in s.tables:
            if t.shape_id in marked:  # таблица с метками заполняется метками
                keep.append(t.shape_id)
                continue
            name = f"table_{s.slide_id}_{t.shape_id}"
            heads = [h.strip() for h in t.header] if t.first_row else []
            if len(heads) != t.cols or len(set(heads)) != t.cols or not all(heads):
                heads = [f"Столбец {j + 1}" for j in range(t.cols)]
            rows = max(1, t.rows - (1 if t.first_row else 0))
            table: dict[str, list[Any]] = {heads[0]: [f"Строка {i + 1}" for i in range(rows)]}
            for h in heads[1:]:
                table[h] = [rnd.randint(1, 100_000) for _ in range(rows)]
            datasets[name] = pa.table(table)
            blocks.append({"type": "table_fill", "shape": t.shape_id, "dataset": name, "header": t.first_row})
        specs.append({"example": s.slide_id, "markers": markers, "blocks": blocks, "keep": keep})
    scenario = ScenarioSpec.model_validate(
        {
            "name": "Приёмочный тест шаблона",
            "inputs": [{"id": "sample", "source": "sample", "main": True}],
            "datasets": [{"id": d, "input": "sample"} for d in datasets],
            "slides": specs,
        }
    )
    engine = EngineResult(period=PERIOD, datasets=datasets)
    return Plan(scenario, engine, slides, plan_values, skipped, categories)


def build(theme: ThemeManifest, plan: Plan, out: Path) -> Built:
    result = build_presentation(plan.scenario, theme, plan.engine, PluginRegistry.discover(), out)
    return Built(plan, result, out)


def errors(result: RenderResult) -> list[str]:
    return [str(i) for i in result.issues if i.level == IssueLevel.ERROR]


@pytest.fixture(scope="module")
def theme(tmp_path_factory: pytest.TempPathFactory) -> ThemeManifest:
    if not TEMPLATE.exists():
        pytest.fail(f"Шаблон не найден: {TEMPLATE}")
    return import_template(TEMPLATE, tmp_path_factory.mktemp("theme"))


@pytest.fixture(scope="module")
def built(theme: ThemeManifest, tmp_path_factory: pytest.TempPathFactory) -> Built:
    if not theme.slides:
        pytest.skip("в шаблоне нет слайдов-образцов")
    b = build(theme, make_plan(theme, SEEDS[0]), tmp_path_factory.mktemp("out") / "report.pptx")
    assert b.result.output_path, "сборка остановилась:\n" + "\n".join(errors(b.result))
    return b


# --- проверки -------------------------------------------------------------------------------


def c14n(el: etree._Element | None, drop: frozenset[str] | set[str] = frozenset()) -> bytes:
    if el is None:
        return b""
    el = copy.deepcopy(el)
    for a in drop:
        el.attrib.pop(a, None)
    return etree.tostring(el, method="c14n", exclusive=True)


def where(s: TemplateSlideInfo, shape: str) -> str:
    return f"слайд {s.number} шаблона, «{shape}»"


def test_unreplaceable_markers(theme: ThemeManifest):
    """Метки, которые приложение заменить не может, правятся в самом шаблоне. В синтетическом
    шаблоне такая метка одна, намеренно: на ней проверяется проверка шаблона."""
    bad = [(s, m) for s in theme.slides for m in s.markers if not m.replaceable]
    expected = {"Метка с переносом"} if TEMPLATE == SYNTHETIC else set()
    assert {m.shape_name for _, m in bad} == expected, "\n".join(
        f"{where(s, m.shape_name)}: {{{{{m.name}}}}} — {m.reason}" for s, m in bad
    )


def test_markers_replaced_with_template_formatting(theme: ThemeManifest, built: Built):
    plan = built.plan
    tpl = template_slides(Presentation(theme.pptx_path))
    want: dict[tuple[int, int, str, int], bytes] = {}
    for s in plan.slides:
        tree = tpl[s.slide_id]._element.find(f"{P}cSld/{P}spTree")
        for sh in iter_shapes(tree):
            for m in iter_markers(sh.el):
                if m.reason is None and m.covered:
                    rpr = m.covered[0].el.find(f"{A}rPr")
                    want[(s.slide_id, sh.id, m.name, m.occurrence)] = c14n(rpr, FREE_RUN_ATTRS)
    out = Presentation(str(built.path))
    found: Counter[str] = Counter()
    problems: list[str] = []
    for s, slide in zip(plan.slides, out.slides, strict=True):
        tree = slide._element.find(f"{P}cSld/{P}spTree")
        for sh in iter_shapes(tree):
            if sh.el.tag != f"{P}grpSp" and "{{" in all_text(sh.el):
                problems.append(f"{where(s, sh.name)}: осталось «{{{{»")
        for r in tree.iter(f"{A}r"):
            value = r.findtext(f"{A}t") or ""
            if value not in plan.values:
                continue
            found[value] += 1
            ts, m = plan.values[value]
            if ts is not s:
                problems.append(f"{where(ts, m.shape_name)}: значение {{{{{m.name}}}}} попало на слайд {s.number}")
                continue
            got = c14n(r.find(f"{A}rPr"), FREE_RUN_ATTRS)
            was = want.get((s.slide_id, m.shape_id, m.name, m.occurrence))
            if got != was:
                problems.append(
                    f"{where(s, m.shape_name)}: оформление значения {{{{{m.name}}}}} не как у метки:\n"
                    f"  было  {(was or b'').decode()}\n  стало {got.decode()}"
                )
    for value, (s, m) in plan.values.items():
        if found[value] != 1:
            problems.append(f"{where(s, m.shape_name)}: значение {{{{{m.name}}}}} в отчёте {found[value]} раз(а)")
    assert not problems, "\n".join(problems)


def _rels(z: zipfile.ZipFile) -> dict[str, set[str]]:
    """Часть → части, которые на неё ссылаются."""
    refs: dict[str, set[str]] = {}
    for name in z.namelist():
        if not name.endswith(".rels"):
            continue
        folder, base = posixpath.split(posixpath.dirname(name))[0], posixpath.basename(name)[: -len(".rels")]
        source = posixpath.join(folder, base)
        for rel in etree.fromstring(z.read(name)):
            if rel.get("TargetMode") == "External":
                continue
            target = posixpath.normpath(posixpath.join(folder, rel.get("Target", ""))).lstrip("/")
            refs.setdefault(target, set()).add(source)
    return refs


def _binary(name: str) -> bool:
    return not name.endswith((".xml", ".rels")) and not name.endswith("/")


def test_binary_parts_and_unfilled_xml_unchanged(theme: ThemeManifest, built: Built):
    problems: list[str] = []
    with zipfile.ZipFile(theme.pptx_path) as zt, zipfile.ZipFile(built.path) as zo:
        tpl_names, out_names = set(zt.namelist()), set(zo.namelist())
        tpl_refs, out_refs = _rels(zt), _rels(zo)
        tpl_hashes = {hashlib.sha256(zt.read(n)).hexdigest() for n in tpl_names if _binary(n)}
        for n in sorted(out_names):
            if not _binary(n) or any(s.startswith("ppt/charts/") for s in out_refs.get(n, ())):
                continue  # книги Excel графиков переписываются данными
            if hashlib.sha256(zo.read(n)).hexdigest() not in tpl_hashes:
                problems.append(f"{n}: такой двоичной части в шаблоне не было")
        for n in sorted(tpl_names):
            shared = any(s.startswith(NOT_SLIDE_OWNED) or s == "ppt/presentation.xml" for s in tpl_refs.get(n, ()))
            if _binary(n) and shared and (n not in out_names or zo.read(n) != zt.read(n)):
                problems.append(f"{n}: шрифт или картинка мастера изменились или пропали")
            if n.startswith(NOT_SLIDE_OWNED) and n.endswith(".xml"):
                if n not in out_names:
                    problems.append(f"{n}: пропала")
                elif c14n(etree.fromstring(zt.read(n))) != c14n(etree.fromstring(zo.read(n))):
                    problems.append(f"{n}: изменилась")

    tpl = template_slides(Presentation(theme.pptx_path))
    out = Presentation(str(built.path))
    for s, slide in zip(built.plan.slides, out.slides, strict=True):
        before, after = tpl[s.slide_id], slide
        pics = [
            Counter(
                hashlib.sha256(r.target_part.blob).hexdigest()
                for r in sl.part.rels.values()
                if not r.is_external and not isinstance(r.target_part, XmlPart)
            )
            for sl in (before, after)
        ]
        if pics[0] != pics[1]:
            problems.append(f"слайд {s.number} шаблона: картинки и вложения слайда изменились")
        filled = {m.shape_id for m in s.markers} | {c.shape_id for c in s.charts} | {t.shape_id for t in s.tables}
        now = {sh.id: sh for sh in iter_shapes(after._element.find(f"{P}cSld/{P}spTree"))}
        for sh in iter_shapes(before._element.find(f"{P}cSld/{P}spTree")):
            if sh.el.tag == f"{P}grpSp" or sh.id in filled:
                continue
            empty_ph = sh.el.find(f".//{P}nvPr/{P}ph") is not None and not all_text(sh.el).strip()
            if sh.id not in now:
                if not empty_ph:  # пустые плейсхолдеры удаляются намеренно
                    problems.append(f"{where(s, sh.name)}: фигура пропала")
            elif c14n(sh.el) != c14n(now[sh.id].el):
                problems.append(f"{where(s, sh.name)}: незаполняемая фигура изменилась")
    assert not problems, "\n".join(problems)


def _chart_problems(theme: ThemeManifest, b: Built) -> list[str]:
    problems: list[str] = []
    out = Presentation(str(b.path))
    for s, slide in zip(b.plan.slides, out.slides, strict=True):
        if not s.charts:
            continue
        shapes = shapes_by_id(slide)
        for c in s.charts:
            cs = shapes[c.shape_id].chart._chartSpace
            counts = [len(group_series(g)) for g in groups(cs)]
            if counts != c.series_count:
                problems.append(f"{where(s, c.shape_name)}: серий по группам {counts}, в шаблоне {c.series_count}")
            for p in check_chart(cs, b.plan.categories[(s.slide_id, c.shape_id)]):
                problems.append(f"{where(s, c.shape_name)}: {p}")
        for p in label_overlaps(slide):
            a, z = ({**x, "chart": x["chart"].split(":", 1)[-1]} for x in (p["a"], p["b"]))
            problems.append(
                f"слайд {s.number} шаблона: подпись «{a['text']}» («{a['chart']}», {a['series']}) "
                f"ложится на «{z['text']}» («{z['chart']}», {z['series']})"
            )
    return problems


def test_charts_keep_series_axes_and_labels_apart(theme: ThemeManifest, built: Built, tmp_path: Path):
    problems = [f"данные {SEEDS[0]}: {p}" for p in _chart_problems(theme, built)]
    for seed in SEEDS[1:]:
        b = build(theme, make_plan(theme, seed), tmp_path / f"report_{seed}.pptx")
        if not b.result.output_path:
            problems += [f"данные {seed}: сборка остановилась: {e}" for e in errors(b.result)]
            continue
        problems += [f"данные {seed}: {p}" for p in _chart_problems(theme, b)]
    assert not problems, "\n".join(problems)


def test_skipped_slide_removed(built: Built):
    if built.plan.skipped is None:
        pytest.skip("в шаблоне один слайд-образец")
    out = Presentation(str(built.path))
    ids = [int(e.get("id")) for e in out.slides._sldIdLst]
    assert len(ids) == len(built.plan.slides)
    assert built.plan.skipped.slide_id not in ids
    with zipfile.ZipFile(built.path) as z:
        parts = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    assert len(parts) == len(built.plan.slides), parts


def test_missing_binding_stops_build_with_places(theme: ThemeManifest, built: Built, tmp_path: Path):
    plan = built.plan
    spec = plan.scenario.model_copy(deep=True)
    marker = next(((k, s) for k, s in enumerate(plan.slides) if s.markers), None)
    chart = next(((k, s) for k, s in enumerate(plan.slides) if s.charts), None)
    if marker is None and chart is None:
        pytest.skip("на слайдах-образцах нет ни меток, ни графиков")
    expect: list[tuple[int, str]] = []
    if marker is not None:
        k, s = marker
        m = s.markers[0]
        del spec.slides[k].markers[m.key]
        expect.append((k + 1, f"«{m.shape_name}» (id {m.shape_id})"))
    if chart is not None:
        k, s = chart
        c = s.charts[0]
        spec.slides[k].blocks = [b for b in spec.slides[k].blocks if b.shape is None or b.shape.id != c.shape_id]
        expect.append((k + 1, f"«{c.shape_name}» (id {c.shape_id})"))
    res = build_presentation(spec, theme, plan.engine, PluginRegistry.discover(), tmp_path / "x.pptx")
    assert res.output_path is None
    text = "\n".join(errors(res))
    for number, shape in expect:
        assert any(f"slide:{number}/" in e and shape in e for e in errors(res)), (number, shape, text)


def _has_powerpoint() -> bool:
    if sys.platform != "win32" or shutil.which("powershell") is None:
        return False
    import winreg

    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application"))
    except OSError:
        return False
    return True


@pytest.mark.skipif(not _has_powerpoint(), reason="нужен PowerPoint")
def test_opens_in_powerpoint(built: Built, tmp_path: Path):
    from autogenerator.worker.theme_jobs import _powerpoint

    assert _powerpoint(built.path, tmp_path / "slide.png"), "PowerPoint не открыл собранный файл"
