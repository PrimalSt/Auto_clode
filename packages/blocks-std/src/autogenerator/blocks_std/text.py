"""Текстовый блок: текст с переменными на Jinja2 в песочнице (ARCHITECTURE.md, раздел 6.5).

Контекст: ``metrics`` (показатели по id), ``period`` (поля отчётного периода), ``scenario``.
Фильтры: ``number``, ``percent``, ``money``, ``date``, ``month``; функции ``change`` и
``change_pct``. Пустой показатель выводится как «нет данных».

Пример::

    Выручка {{ metrics.revenue | money(scale='million', decimals=1) }}
    ({{ change_pct(metrics.revenue, metrics.revenue_prev) | percent(sign=true) }} к прошлому месяцу)
"""

from __future__ import annotations

from typing import Any, Literal

from babel.numbers import format_decimal
from jinja2 import StrictUndefined, TemplateSyntaxError, UndefinedError, nodes
from jinja2.sandbox import SandboxedEnvironment
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt
from pydantic import BaseModel, ConfigDict, Field

from autogenerator.contracts import BlockContext, BlockData, BlockPlugin, BlockTarget, DataNeeds

from .formats import (
    LOCALE,
    NO_DATA,
    PeriodVars,
    fmt_date,
    fmt_money,
    fmt_number,
    fmt_percent,
    month_name,
)


def change(a: Any, b: Any) -> float | None:
    """Разница ``a − b``; пусто, если одно из значений пустое."""
    if a is None or b is None:
        return None
    return float(a) - float(b)


def change_pct(a: Any, b: Any) -> float | None:
    """Относительное изменение ``a / b − 1``; пусто при пустом значении или нулевой базе."""
    if a is None or b is None or float(b) == 0:
        return None
    return float(a) / float(b) - 1


def _finalize(value: Any) -> Any:
    if value is None:
        return NO_DATA
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, int | float):
        return format_decimal(value, format="#,##0.##", locale=LOCALE)
    return value


def _month_filter(value: Any, case: str = "nom", capital: bool = False) -> str | None:
    if value is None:
        return None
    return month_name(value, case, capital)  # type: ignore[arg-type]


def make_environment() -> SandboxedEnvironment:
    env = SandboxedEnvironment(
        autoescape=False,
        undefined=StrictUndefined,
        finalize=_finalize,
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.filters.update(
        number=fmt_number,
        percent=fmt_percent,
        money=fmt_money,
        date=fmt_date,
        month=_month_filter,
    )
    env.globals.update(change=change, change_pct=change_pct)
    return env


def referenced_metrics(text: str) -> set[str]:
    """id показателей, на которые ссылается текст (``metrics.x`` и ``metrics['x']``)."""
    ast = make_environment().parse(text)
    found: set[str] = set()
    for n in ast.find_all(nodes.Getattr):
        if isinstance(n.node, nodes.Name) and n.node.name == "metrics":
            found.add(n.attr)
    for g in ast.find_all(nodes.Getitem):
        if isinstance(g.node, nodes.Name) and g.node.name == "metrics" and isinstance(g.arg, nodes.Const):
            found.add(str(g.arg.value))
    return found


def render_text(text: str, metrics: dict[str, Any], ctx: BlockContext) -> str:
    env = make_environment()
    try:
        tpl = env.from_string(text)
        return tpl.render(
            metrics=metrics,
            period=PeriodVars(ctx.period),
            scenario={"name": ctx.scenario_name},
        )
    except TemplateSyntaxError as e:
        raise ValueError(f"ошибка в тексте (строка {e.lineno}): {e.message}") from e
    except UndefinedError as e:
        raise ValueError(f"в тексте неизвестная переменная: {e.message}") from e


class TextParams(BaseModel):
    model_config = ConfigDict(extra="forbid")

    text: str = Field(description="Текст; переменные — в {{ … }}")
    font_size: float | None = Field(None, gt=0, description="Размер шрифта, пт; пусто — как в макете")
    bold: bool | None = None
    align: Literal["left", "center", "right"] | None = None


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


class TextBlock(BlockPlugin):
    name = "text"
    title = "Текст"
    Params = TextParams

    def data_needs(self, params: Any) -> DataNeeds:
        try:
            return DataNeeds(metrics=referenced_metrics(params.text))
        except TemplateSyntaxError as e:
            raise ValueError(f"ошибка в тексте (строка {e.lineno}): {e.message}") from e

    def render(self, target: BlockTarget, params: Any, data: BlockData, ctx: BlockContext) -> None:
        text = render_text(params.text, data.metrics, ctx)
        ph = target.placeholder
        if ph is not None and getattr(ph, "has_text_frame", False):
            tf = ph.text_frame
        else:
            g = target.geometry
            box = target.slide.shapes.add_textbox(Emu(g.x), Emu(g.y), Emu(g.cx), Emu(g.cy))
            tf = box.text_frame
            tf.word_wrap = True
        lines = text.split("\n")
        tf.text = lines[0]
        for line in lines[1:]:
            tf.add_paragraph().text = line
        for p in tf.paragraphs:
            if params.align:
                p.alignment = _ALIGN[params.align]
            for r in p.runs:
                if params.font_size:
                    r.font.size = Pt(params.font_size)
                if params.bold is not None:
                    r.font.bold = params.bold
