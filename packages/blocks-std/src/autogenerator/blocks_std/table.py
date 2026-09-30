"""Блок «таблица»: новая таблица PowerPoint на слайде из макета, с русскими форматами чисел
и дат и усечением до ``max_rows`` строк."""

from __future__ import annotations

from datetime import date, datetime
from typing import Any, Literal

from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt
from pydantic import BaseModel, ConfigDict, Field, model_validator

from autogenerator.contracts import BlockContext, BlockData, BlockPlugin, BlockTarget, DataNeeds

from .formats import fmt_date, fmt_money, fmt_number, fmt_percent

# Пустое значение в ячейке: прочерк, а не пустота — видно, что данных нет, а не что их забыли.
EMPTY_CELL = "—"


class CellFormat(BaseModel):
    model_config = ConfigDict(extra="forbid")

    kind: Literal["auto", "number", "percent", "money", "date", "text"] = "auto"
    decimals: int | None = Field(None, ge=0, le=6)
    scale: Literal["thousand", "million", "billion"] | None = None
    sign: bool = False
    pattern: str | None = Field(None, description="Для дат: шаблон Babel, например LLLL yyyy")


class TableColumn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    column: str
    header: str | None = None
    format: CellFormat = Field(default_factory=CellFormat)

    @model_validator(mode="before")
    @classmethod
    def _from_text(cls, v: Any) -> Any:
        return {"column": v} if isinstance(v, str) else v


class TableParams(BaseModel):
    model_config = ConfigDict(extra="forbid")

    dataset: str
    columns: list[TableColumn] = Field(default_factory=list, description="Пусто — все столбцы набора")
    max_rows: int = Field(20, ge=1)
    font_size: float = Field(12, gt=0)


def _auto_decimals(values: list[Any]) -> int:
    nums = [float(v) for v in values if isinstance(v, int | float) and not isinstance(v, bool)]
    if not nums or all(n.is_integer() for n in nums) or max(abs(n) for n in nums) >= 100:
        return 0
    return 2


def format_column(values: list[Any], fmt: CellFormat) -> tuple[list[str], bool]:
    """Тексты ячеек столбца и признак «числовой» (для выравнивания вправо)."""
    kind = fmt.kind
    sample = next((v for v in values if v is not None), None)
    if kind == "auto":
        if isinstance(sample, bool):
            kind = "text"
        elif isinstance(sample, int | float):
            kind = "number"
        elif isinstance(sample, datetime | date):
            kind = "date"
        else:
            kind = "text"
    out: list[str] = []
    if kind == "number":
        d = fmt.decimals if fmt.decimals is not None else _auto_decimals(values)
        out = [fmt_number(v, decimals=d, scale=fmt.scale, sign=fmt.sign) or EMPTY_CELL for v in values]
    elif kind == "percent":
        d = fmt.decimals if fmt.decimals is not None else 1
        out = [fmt_percent(v, decimals=d, sign=fmt.sign) or EMPTY_CELL for v in values]
    elif kind == "money":
        d = fmt.decimals if fmt.decimals is not None else 0
        out = [fmt_money(v, decimals=d, scale=fmt.scale) or EMPTY_CELL for v in values]
    elif kind == "date":
        dates = [v.date() if isinstance(v, datetime) else v for v in values if v is not None]
        # Все даты — первые числа месяцев: это группировка по месяцам, показываем «январь 2026».
        default = "LLLL yyyy" if dates and all(getattr(d, "day", 0) == 1 for d in dates) else "dd.MM.yyyy"
        out = [fmt_date(v, fmt.pattern or default) or EMPTY_CELL for v in values]
    else:
        out = [EMPTY_CELL if v is None else str(v) for v in values]
    return out, kind in ("number", "percent", "money")


class TableBlock(BlockPlugin):
    name = "table"
    title = "Таблица"
    Params = TableParams

    def data_needs(self, params: Any) -> DataNeeds:
        return DataNeeds(datasets={params.dataset})

    def render(self, target: BlockTarget, params: Any, data: BlockData, ctx: BlockContext) -> None:
        table = data.datasets[params.dataset]
        cols = params.columns or [TableColumn(column=c) for c in table.column_names]
        for c in cols:
            if c.column not in table.column_names:
                raise ValueError(
                    f"в наборе «{params.dataset}» нет столбца «{c.column}» (есть: {', '.join(table.column_names)})"
                )
        n = table.num_rows
        if n > params.max_rows:
            ctx.warn(f"в таблице {n} строк, показаны первые {params.max_rows}")
            table = table.slice(0, params.max_rows)
            n = params.max_rows
        cells = [format_column(table.column(c.column).to_pylist(), c.format) for c in cols]

        g = target.geometry
        row_h = min(g.cy // (n + 1), int(Pt(params.font_size * 2.2)))
        frame = target.slide.shapes.add_table(n + 1, len(cols), Emu(g.x), Emu(g.y), Emu(g.cx), Emu(row_h * (n + 1)))
        tbl = frame.table
        for r in range(n + 1):
            tbl.rows[r].height = Emu(row_h)
        for j, (c, (texts, numeric)) in enumerate(zip(cols, cells, strict=True)):
            self._cell(tbl.cell(0, j), c.header or c.column, params.font_size, numeric, bold=True)
            for i, t in enumerate(texts, start=1):
                self._cell(tbl.cell(i, j), t, params.font_size, numeric)

    @staticmethod
    def _cell(cell: Any, text: str, size: float, numeric: bool, bold: bool = False) -> None:
        cell.text = text
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(size)
            if bold:
                r.font.bold = True
