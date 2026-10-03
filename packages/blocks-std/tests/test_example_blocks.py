"""Блоки слайдов-образцов на синтетическом шаблоне examples/templates/synthetic.pptx."""

from dataclasses import dataclass, field
from pathlib import Path

import pyarrow as pa
import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from autogenerator.blocks_std.chart_fill import ChartFillBlock
from autogenerator.blocks_std.formats import NBSP
from autogenerator.blocks_std.markers import MarkersBlock
from autogenerator.blocks_std.table_fill import TableFillBlock
from autogenerator.contracts import BlockData, BlockError, BlockTarget, Geometry, Period
from autogenerator.contracts.ooxml import A, C, all_text
from autogenerator.theme.slides import slide_info

TEMPLATE = Path(__file__).parents[3] / "examples" / "templates" / "synthetic.pptx"


@dataclass
class Ctx:
    period: Period = field(default_factory=lambda: Period.parse("2026-03"))
    scenario_name: str = "Тест"
    preview: bool = False
    warnings: list[str] = field(default_factory=list)
    kept: list[tuple[int, str]] = field(default_factory=list)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def keep_marker(self, shape_id: int, name: str) -> None:
        self.kept.append((shape_id, name))


def example(number: int, prs=None):
    prs = prs or Presentation(str(TEMPLATE))
    slide = prs.slides[number - 1]
    return prs, slide, slide_info(slide, number, slide.slide_id, "")


def shape_target(slide, info, shape_id: int) -> BlockTarget:
    shape = next(s for s in slide.shapes if s.shape_id == shape_id)
    return BlockTarget(slide=slide, geometry=Geometry(x=0, y=0, cx=1, cy=1), shape=shape, example=info)


def slide_target(slide, info) -> BlockTarget:
    return BlockTarget(slide=slide, geometry=Geometry(x=0, y=0, cx=1, cy=1), example=info)


def text_slide(*paragraphs: list[str]):
    """Слайд с одной надписью: каждый абзац — список прогонов (метка разрезана, как в PowerPoint)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    tf = box.text_frame
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for j, text in enumerate(runs):
            r = p.add_run()
            r.text = text
            r.font.size = Pt(20 + j)
            r.font.bold = j == 1
    return slide, box, slide_info(slide, 1, slide.slide_id, "")


def run(block, target, params: dict, data: BlockData | None = None, ctx: Ctx | None = None) -> Ctx:
    ctx = ctx or Ctx()
    block.render(target, block.parse_params(params), data or BlockData(), ctx)
    return ctx


# --- метки --------------------------------------------------------------------------------

COMMON = {"Месяц": "period.month", "Год": "period.year"}


def test_markers_fill_example_slide():
    _, slide, info = example(2)
    assert set(info.marker_names) == {"Месяц", "Год", "Выручка", "Прирост+", "Доля"}
    params = {
        "bindings": {
            "Выручка": {"metric": "revenue", "scale": "млн", "decimals": 1},
            "{{Прирост+}}": {"metric": "growth", "percent": True, "sign": True},
            "Доля": {"metric": "share", "percent": True, "decimals": 0},
        },
        "common": COMMON,
    }
    block = MarkersBlock()
    assert block.check(block.parse_params(params), info, None) == []
    data = BlockData(metrics={"revenue": 26_229_878, "growth": 0.125, "share": 0.4})
    run(block, slide_target(slide, info), params, data)
    text = "\n".join(all_text(s._element) for s in slide.shapes)
    assert "{{" not in text
    assert "Итоги, март 2026" in text
    assert "26,2" in text and "+12,5%" in text and "40%" in text


def test_marker_split_into_runs_keeps_formatting():
    slide, box, info = text_slide(["Рост за ", "{{Ме", "сяц}}", " — {{Доля}} %"])
    block = MarkersBlock()
    run(
        block,
        slide_target(slide, info),
        {"bindings": {"Доля": {"metric": "x", "percent": True}}, "common": {"Месяц": "period.month_prep"}},
        BlockData(metrics={"x": 0.0567}),
    )
    p = box.text_frame.paragraphs[0]
    # «%» после метки в шаблоне: значение подставляется без своего знака
    assert p.text == "Рост за марте — 5,7 %"
    runs = p.runs
    value = next(r for r in runs if r.text == "марте")
    assert value.font.size == Pt(21) and value.font.bold  # оформление прогона, где начиналась метка
    assert value._r.find(f"{A}rPr").get("lang") == "ru-RU"


def test_marker_binding_precedence():
    slide, box, info = text_slide(["{{X}} и {{X}}"], ["ещё {{X}}"])
    shape = box.shape_id
    params = {
        "bindings": {"X": {"value": "слайд"}, f"X@{shape}#2": {"value": "второе"}},
        "common": {"X": {"value": "общее"}},
    }
    run(MarkersBlock(), slide_target(slide, info), params)
    assert [p.text for p in box.text_frame.paragraphs] == ["слайд и второе", "ещё слайд"]


def test_marker_empty_value_behaviours():
    def go(empty: str, preview: bool = False) -> tuple[str, Ctx]:
        slide, box, info = text_slide(["Итого: {{Сумма}}"])
        ctx = Ctx(preview=preview)
        params = {"bindings": {"Сумма": {"metric": "total", "empty": empty}}}
        run(MarkersBlock(), slide_target(slide, info), params, BlockData(metrics={"total": None}), ctx)
        return box.text_frame.text, ctx

    with pytest.raises(BlockError, match="нет значения"):
        go("stop")
    assert go("blank")[0] == "Итого: "
    text, ctx = go("keep")
    assert text == "Итого: {{Сумма}}" and ctx.kept and not ctx.warnings
    text, ctx = go("stop", preview=True)
    assert text == "Итого: {{Сумма}}" and ctx.kept and ctx.warnings


def test_markers_check_reports_unbound_and_stray_bindings():
    _, _, info = example(2)
    block = MarkersBlock()
    problems = block.check(block.parse_params({"bindings": {"Нет такой": "x"}, "common": COMMON}), info, None)
    assert any("{{Выручка}} не привязана" in p for p in problems)
    assert any("нет метки для привязки «Нет такой»" in p for p in problems)
    with pytest.raises(ValueError, match="без фигуры"):
        block.parse_params({"common": {"X@5": "x"}})


def test_unreplaceable_marker_stops_build():
    _, slide, info = example(7)
    block = MarkersBlock()
    params = block.parse_params({"common": COMMON})
    problems = block.check(params, info, None)
    assert any("перенос строки" in p for p in problems)
    with pytest.raises(BlockError):
        block.render(slide_target(slide, info), params, BlockData(), Ctx())


def run_color(r) -> str | None:
    clr = r._r.find(f"{A}rPr/{A}solidFill/{A}srgbClr")
    return None if clr is None else clr.get("val")


def colored(binding: dict, metrics: dict | None = None) -> dict[str, str | None]:
    """Цвета прогонов абзаца «Δ {{X}} к марту» после замены: текст прогона → RGB."""
    slide, box, info = text_slide(["Δ {{X}} к марту"])
    run(MarkersBlock(), slide_target(slide, info), {"bindings": {"X": binding}}, BlockData(metrics=metrics or {}))
    return {r.text: run_color(r) for r in box.text_frame.paragraphs[0].runs}


def test_marker_color_forms():
    def color(value):
        return MarkersBlock().parse_params({"bindings": {"X": {"value": "1", "color": value}}}).bindings["X"].color

    assert MarkersBlock().parse_params({"bindings": {"X": "x"}}).bindings["X"].color is None
    c = color("sign")
    assert (c.positive, c.negative, c.zero, c.by) == ("2E7559", "C00000", None, None)
    c = color("1f4e79")
    assert (c.positive, c.negative, c.zero) == ("1F4E79",) * 3
    c = color({"positive": "#00b050", "zero": "7f7f7f", "negative": None, "by": "d"})
    assert (c.positive, c.negative, c.zero, c.by) == ("00B050", None, "7F7F7F", "d")
    for bad in ("12345", "GGGGGG", {"negative": "красный"}):
        with pytest.raises(ValueError, match="6 шестнадцатеричных цифр"):
            color(bad)
    for number in (123456, float("inf")):  # 123456 и 2E7559 без кавычек в YAML
        with pytest.raises(ValueError, match="в кавычки"):
            color(number)
    with pytest.raises(ValueError, match="extra"):
        color({"positiv": "2E7559"})


def test_marker_color_by_sign():
    by_metric = {"metric": "d", "percent": True, "color": "sign"}
    assert colored(by_metric, {"d": 0.05}) == {"Δ ": None, "5,0%": "2E7559", " к марту": None}
    assert colored(by_metric, {"d": -0.05})["-5,0%"] == "C00000"
    assert colored(by_metric, {"d": 0})["0,0%"] is None
    # by важнее показателя привязки и текста; пустой by — знак из текста, а не из показателя привязки
    assert colored({"metric": "d", "color": {"by": "e"}}, {"d": -3, "e": 1})["-3"] == "2E7559"
    assert colored({"value": "-1", "color": {"by": "e"}}, {"e": 0.5})["-1"] == "2E7559"
    assert colored({"value": "(-1%)", "color": {"by": "e"}}, {"e": None})["(-1%)"] == "C00000"
    assert colored({"metric": "d", "color": {"by": "e"}}, {"d": 5, "e": None})["5"] is None
    assert colored({"value": "-1", "color": {"by": "e"}}, {"e": 10**400})["-1"] == "C00000"  # не влезает во float
    for text, rgb in (("(+3%)", "2E7559"), ("(-3%)", "C00000"), ("(\u22123%)", "C00000"), ("(0%)", None)):
        assert colored({"value": text, "color": "sign"})[text] == rgb
    text = {"text": "({{ metrics.d | number(sign=true) }}%)", "color": "sign"}
    assert colored(text, {"d": 3})["(+3%)"] == "2E7559"
    assert colored(text, {"d": -2})["(-2%)"] == "C00000"
    fixed = {"value": "(-1%)", "color": "1F4E79"}
    assert colored(fixed)["(-1%)"] == "1F4E79"
    zero = {"value": "0", "color": {"zero": "7F7F7F", "positive": None}}
    assert colored(zero)["0"] == "7F7F7F"
    assert colored({"value": "+1", "color": {"positive": None}})["+1"] is None


def test_marker_color_replaces_template_fill_in_schema_order():
    slide, box, info = text_slide(["{{X}}", " и {{Y}}"])
    for r in box.text_frame.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        r.font.name = "Test Sans"
        r._r.find(f"{A}rPr").insert(0, etree.Element(f"{A}ln"))
    params = {"bindings": {"X": {"value": "+5", "color": "sign"}, "Y": {"value": "0", "color": "sign"}}}
    run(MarkersBlock(), slide_target(slide, info), params)
    runs = {r.text: r._r.find(f"{A}rPr") for r in box.text_frame.paragraphs[0].runs}
    assert [etree.QName(e).localname for e in runs["+5"]] == ["ln", "solidFill", "latin"]
    assert runs["+5"].find(f"{A}solidFill/{A}srgbClr").get("val") == "2E7559"
    # ноль без цвета — цвет шаблона остаётся
    assert runs["0"].find(f"{A}solidFill/{A}srgbClr").get("val") == "123456"
    assert runs[" и "].find(f"{A}solidFill/{A}srgbClr").get("val") == "123456"

    # без заливки в шаблоне: перед эффектами и шрифтом; другая заливка (gradFill) убирается
    for children, expected in (
        (["ln", "effectLst", "latin"], ["ln", "solidFill", "effectLst", "latin"]),
        (["gradFill", "highlight", "latin"], ["solidFill", "highlight", "latin"]),
    ):
        slide, box, info = text_slide(["{{X}}"])
        rpr = box.text_frame.paragraphs[0].runs[0]._r.find(f"{A}rPr")
        for tag in children:
            etree.SubElement(rpr, f"{A}{tag}")
        run(MarkersBlock(), slide_target(slide, info), {"bindings": {"X": {"value": "-5", "color": "sign"}}})
        rpr = box.text_frame.paragraphs[0].runs[0]._r.find(f"{A}rPr")
        assert [etree.QName(e).localname for e in rpr] == expected


def test_marker_color_skips_blank_and_preview():
    binding = {"metric": "d", "empty": "blank", "color": {"zero": "7F7F7F"}}
    assert colored(binding, {"d": None}) == {"Δ ": None, "": None, " к марту": None}
    ctx = Ctx(preview=True)
    slide, box, info = text_slide(["{{X}}"])
    params = {"bindings": {"X": {"metric": "d", "color": {"zero": "7F7F7F"}}}}
    run(MarkersBlock(), slide_target(slide, info), params, BlockData(metrics={"d": None}), ctx)
    rpr = box.text_frame.paragraphs[0].runs[0]._r.find(f"{A}rPr")
    assert rpr.find(f"{A}highlight") is not None and rpr.find(f"{A}solidFill") is None


def test_marker_color_data_needs():
    block = MarkersBlock()
    params = {
        "bindings": {
            "A": {"value": "x", "color": {"by": "a"}},
            "B": {"metric": "b", "color": "sign"},
            "C": {"text": "{{ metrics.c }}", "color": {"by": "c2"}},
        },
        "common": {"D": {"period": "month", "color": {"by": "d"}}},
    }
    assert block.data_needs(block.parse_params(params)).metrics == {"a", "b", "c", "c2", "d"}


# --- графики ------------------------------------------------------------------------------


def chart_xml(slide, shape_id: int):
    shape = next(s for s in slide.shapes if s.shape_id == shape_id)
    return shape.chart._chartSpace


def series_by_group(cs) -> list[int]:
    plot = cs.find(f"{C}chart/{C}plotArea")
    return [len(g.findall(f"{C}ser")) for g in plot if g.tag.endswith("Chart")]


MONTHS = pa.table(
    {
        "month": ["янв", "фев", "мар", "апр"],
        "a": [10.0, 12.0, 14.0, 13.0],
        "b": [5.0, 6.0, 5.5, 7.0],
        "c": [3.0, 2.0, 4.0, 5.0],
        "share": [0.31, 0.35, 0.33, 0.36],
    }
)


def test_chart_fill_combo_changes_categories_and_series():
    _, slide, info = example(2)
    block = ChartFillBlock()
    params = {
        "dataset": "m",
        "categories": "month",
        "series": ["a", "b", "c", {"column": "share", "name": "Доля", "number_format": "0%"}],
        "groups": [3, 1],
    }
    assert block.check(block.parse_params(params), info, 6) == []
    t = shape_target(slide, info, 6)
    run(block, t, params, BlockData(datasets={"m": MONTHS}))
    cs = chart_xml(slide, 6)
    assert series_by_group(cs) == [3, 1]
    idx = [int(s.get("val")) for s in cs.iter(f"{C}idx") if s.getparent().tag == f"{C}ser"]
    assert len(set(idx)) == 4
    chart = t.shape.chart
    assert list(chart.plots[0].categories) == ["янв", "фев", "мар", "апр"]
    assert [s.name for p in chart.plots for s in p.series] == ["a", "b", "c", "Доля"]
    colors = [
        s.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val")
        for s in cs.iter(f"{C}ser")
        if s.find(f"{C}spPr/{A}solidFill/{A}srgbClr") is not None
    ]
    assert len(colors) == len(set(colors))
    # Формат из сценария получают и подписи: здесь общие подписи группы линии (одна серия);
    # у столбцов формат задан не у всех серий — остаётся формат шаблона.
    label_formats = [g.find(f"{C}dLbls/{C}numFmt").get("formatCode") for g in cs.iter(f"{C}barChart", f"{C}lineChart")]
    assert label_formats == ["#,##0.0", "0%"]
    params["series"][3]["number_format"] = "0.0%"
    run(block, t, params, BlockData(datasets={"m": MONTHS}))
    assert chart_xml(slide, 6).find(f".//{C}lineChart/{C}dLbls/{C}numFmt").get("formatCode") == "0.0%"
    ctx = Ctx()
    block.finish_slide(slide, [(t, block.parse_params(params))], ctx)


def test_chart_fill_series_count_needs_groups():
    _, _, info = example(2)
    block = ChartFillBlock()
    problems = block.check(block.parse_params({"dataset": "m", "categories": "month", "series": ["a", "b"]}), info, 6)
    assert problems and "groups" in problems[0]
    assert (
        "нет графика с id 99"
        in block.check(block.parse_params({"dataset": "m", "categories": "month", "series": ["a"]}), info, 99)[0]
    )


def test_pie_collects_others():
    _, slide, info = example(3)
    regions = pa.table({"r": [f"Регион {i}" for i in range(9)], "v": [float(9 - i) for i in range(9)]})
    t = shape_target(slide, info, 4)
    ctx = run(
        ChartFillBlock(), t, {"dataset": "r", "categories": "r", "series": ["v"]}, BlockData(datasets={"r": regions})
    )
    cats = list(t.shape.chart.plots[0].categories)
    assert len(cats) == 6 and cats[-1] == "Прочие"
    assert list(t.shape.chart.plots[0].series[0].values)[-1] == 1 + 2 + 3 + 4
    assert any("Прочие" in w for w in ctx.warnings)


def test_labels_per_category_fix_category_count():
    _, slide, info = example(4)
    assert info.chart(3).labels_per_category
    four = pa.table({"k": ["а", "б", "в", "г"], "x": [0.5] * 4, "y": [0.5] * 4})
    with pytest.raises(BlockError, match="закреплены"):
        run(
            ChartFillBlock(),
            shape_target(slide, info, 3),
            {"dataset": "d", "categories": "k", "series": ["x", "y"]},
            BlockData(datasets={"d": four}),
        )
    five = pa.table({"k": list("абвгд"), "x": [0.6] * 5, "y": [0.4] * 5})
    run(
        ChartFillBlock(),
        shape_target(slide, info, 3),
        {"dataset": "d", "categories": "k", "series": ["x", "y"]},
        BlockData(datasets={"d": five}),
    )


def test_overlaid_charts_need_same_categories():
    _, slide, info = example(5)
    block = ChartFillBlock()
    a = pa.table({"k": ["1", "2", "3"], "x": [1.0, 2, 3], "y": [1.0, 1, 1], "z": [0.1, 0.2, 0.3]})
    b = pa.table({"k": ["1", "2", "4"], "x": [1.0, 2, 3], "y": [1.0, 1, 1], "z": [0.1, 0.2, 0.3]})
    data = BlockData(datasets={"a": a, "b": b})
    p5 = {"dataset": "a", "categories": "k", "series": ["x", "y"]}
    p6 = {"dataset": "b", "categories": "k", "series": ["x", "y", "z"]}
    t5, t6 = shape_target(slide, info, 5), shape_target(slide, info, 6)
    run(block, t5, p5, data)
    run(block, t6, p6, data)
    with pytest.raises(BlockError, match="наложены"):
        block.finish_slide(slide, [(t5, block.parse_params(p5)), (t6, block.parse_params(p6))], Ctx())


# --- таблица ------------------------------------------------------------------------------


def table_shape(slide):
    return next(s for s in slide.shapes if s.has_table)


def test_table_fill_rows_and_columns():
    _, slide, info = example(6)
    shape = table_shape(slide)
    width = shape.width
    grid_before = [c.width for c in shape.table.columns]
    data = pa.table(
        {
            "region": [f"Регион {i}" for i in range(8)],
            "plan": [1_250_000.0 + i for i in range(8)],
            "share": [0.1 * i for i in range(8)],
        }
    )
    params = {
        "dataset": "t",
        "columns": [
            {"column": "region", "header": "Регион"},
            {"column": "plan", "header": "План, млн", "scale": "million", "decimals": 1},
            {"column": "share", "header": "Доля", "percent": True, "decimals": 0},
        ],
    }
    block = TableFillBlock()
    assert block.check(block.parse_params(params), info, shape.shape_id) == []
    ctx = run(block, shape_target(slide, info, shape.shape_id), params, BlockData(datasets={"t": data}))
    tbl = shape.table
    assert len(tbl.rows) == 9 and len(tbl.columns) == 3
    assert [tbl.cell(0, j).text for j in range(3)] == ["Регион", "План, млн", "Доля"]
    assert [tbl.cell(2, j).text for j in range(3)] == ["Регион 1", "1,3", "10%"]
    assert sum(c.width for c in tbl.columns) == pytest.approx(width, abs=3)
    assert tbl.columns[0].width == grid_before[0]  # первый столбец сохраняет ширину
    row_ids = [e.get("val") for e in tbl._tbl.iter("{http://schemas.microsoft.com/office/drawing/2014/main}rowId")]
    assert len(row_ids) == len(set(row_ids))
    assert not ctx.warnings


def test_table_fill_shrinks_then_keeps_last_rows():
    _, slide, info = example(6)
    shape = table_shape(slide)
    data = pa.table({"k": [f"строка {i}" for i in range(80)], "v": [float(i) for i in range(80)]})
    ctx = run(
        TableFillBlock(),
        shape_target(slide, info, shape.shape_id),
        {"dataset": "t", "min_font_size": 9},
        BlockData(datasets={"t": data}),
    )
    tbl = shape.table
    assert len(tbl.rows) < 81
    assert tbl.cell(len(tbl.rows) - 1, 0).text == "строка 79"  # остаются последние строки
    assert any("показаны последние" in w for w in ctx.warnings)
    sizes = {r.get("sz") for r in tbl._tbl.iter(f"{A}rPr")}
    assert sizes == {"900"}
    assert shape.top + shape.height <= slide.part.package.presentation_part.presentation.slide_height


def test_table_fill_unknown_column():
    _, slide, info = example(6)
    shape = table_shape(slide)
    with pytest.raises(BlockError, match="нет столбцов nope"):
        run(
            TableFillBlock(),
            shape_target(slide, info, shape.shape_id),
            {"dataset": "t", "columns": ["nope"]},
            BlockData(datasets={"t": pa.table({"a": [1]})}),
        )


def test_number_nbsp_in_table_cells():
    _, slide, info = example(6)
    shape = table_shape(slide)
    run(
        TableFillBlock(),
        shape_target(slide, info, shape.shape_id),
        {"dataset": "t"},
        BlockData(datasets={"t": pa.table({"k": ["x"], "v": [1234567.0]})}),
    )
    assert shape.table.cell(1, 1).text == f"1{NBSP}234{NBSP}567"


def fill_table(params: dict, data: pa.Table, prepare=None):
    """Таблица слайда 6 синтетического шаблона после table_fill; ``prepare`` правит слайд до сборки."""
    _, slide, info = example(6)
    shape = table_shape(slide)
    if prepare:
        prepare(slide, shape)
    ctx = run(TableFillBlock(), shape_target(slide, info, shape.shape_id), params, BlockData(datasets={"t": data}))
    return shape, ctx


def fills(shape, column: int) -> list[str | None]:
    """Цвета заливки ячеек столбца в строках данных; ``None`` — своей заливки нет."""
    out = []
    for tr in list(shape.table._tbl.tr_lst)[1:]:
        clr = tr.findall(f"{A}tc")[column].find(f"{A}tcPr/{A}solidFill/{A}srgbClr")
        out.append(clr.get("val") if clr is not None else None)
    return out


def test_table_fill_cell_colors_from_column():
    data = pa.table({"k": ["а", "б", "в", "г"], "v": [1.0, 2, 3, 4], "color": ["ff0000", None, "", "#00FF00"]})
    params = {"dataset": "t", "columns": [{"column": "k", "fill": "color"}], "other_columns": True}
    shape, _ = fill_table(params, data)
    assert [c.text for c in shape.table.rows[0].cells] == ["k", "v"]  # столбец цвета не выводится
    assert fills(shape, 0) == ["FF0000", None, None, "00FF00"]
    assert fills(shape, 1) == [None] * 4
    with pytest.raises(BlockError, match="не цвет RRGGBB: «red»"):
        fill_table(params, data.set_column(2, "color", pa.array(["red", None, None, None])))
    with pytest.raises(BlockError, match="нет столбцов nope"):
        fill_table({"dataset": "t", "columns": [{"column": "k", "fill": "nope"}]}, data)


def test_table_fill_heatmap_colors():
    # Шкала 0…100: красный — жёлтый — зелёный, каналы между цветами — с отбрасыванием дробной части.
    ages = ["0%", "50.0%", "100%", f"25,0{NBSP}%", "-", None, "150", "-10"]
    data = pa.table({"cohort": [f"к{i}" for i in range(len(ages))], "age_1": ages})
    params = {"dataset": "t", "heatmap": {"columns": ["age_1"], "min": 0, "max": 100}}
    shape, _ = fill_table(params, data)
    assert fills(shape, 1) == ["F8696B", "FFEB84", "63BE7B", "FBAA77", None, None, "63BE7B", "F8696B"]
    assert fills(shape, 0) == [None] * len(ages)
    assert shape.table.cell(2, 1).text == "50.0%"  # текст ячейки не меняется


def test_table_fill_heatmap_auto_range_and_explicit_fill_wins():
    data = pa.table(
        {
            "k": ["а", "б", "в"],
            "a1": [20, 30, None],
            "a2": [40, 25, 30],
            "color": ["0000FF", None, None],
        }
    )
    params = {
        "dataset": "t",
        "columns": ["k", {"column": "a1", "fill": "color"}, "a2"],
        "heatmap": {"columns": ["a1", "a2"]},
    }
    shape, _ = fill_table(params, data)
    # Наименьшее (20) и наибольшее (40) — по обоим столбцам карты.
    assert fills(shape, 1) == ["0000FF", "FFEB84", None]
    assert fills(shape, 2) == ["63BE7B", "FBAA77", "FFEB84"]
    with pytest.raises(BlockError, match="heatmap: столбцов color нет в таблице"):
        fill_table({**params, "heatmap": {"columns": ["color"]}}, data)
    with pytest.raises(ValueError, match="min должен быть меньше max"):
        TableFillBlock().parse_params({**params, "heatmap": {"columns": ["a1"], "min": 1, "max": 1}})


def test_table_fill_fill_keeps_schema_order():
    from lxml import etree

    border = f'<a:lnB xmlns:a="{A[1:-1]}" w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:lnB>'

    def prepare(slide, shape):
        trs = list(shape.table._tbl.tr_lst)
        pr = trs[1].findall(f"{A}tc")[0].find(f"{A}tcPr")
        for tag in ("lnL", "lnR"):
            etree.SubElement(pr, f"{A}{tag}", w="6350")
        pr.append(etree.fromstring(border))
        etree.SubElement(pr, f"{A}noFill")
        etree.SubElement(pr, f"{A}extLst")
        tc = trs[2].findall(f"{A}tc")[0]
        tc.remove(tc.find(f"{A}tcPr"))  # у ячейки нет tcPr

    colors = ["AA0000", "00AA00", "0000AA", "AAAA00", "00AAAA"]
    data = pa.table({"k": list("абвгд"), "color": colors})
    shape, _ = fill_table({"dataset": "t", "columns": [{"column": "k", "fill": "color"}]}, data, prepare)
    trs = list(shape.table._tbl.tr_lst)
    for tr in (trs[1], trs[5]):  # строка шаблона и новая, скопированная с неё
        pr = tr.findall(f"{A}tc")[0].find(f"{A}tcPr")
        assert [etree.QName(el).localname for el in pr] == ["lnL", "lnR", "lnB", "solidFill", "extLst"]
        assert pr.find(f"{A}lnB/{A}solidFill/{A}srgbClr").get("val") == "000000"  # граница не тронута
    tc = trs[2].findall(f"{A}tc")[0]
    assert [etree.QName(el).localname for el in tc] == ["txBody", "tcPr"]
    assert fills(shape, 0) == colors


def test_table_fill_colors_follow_trimmed_rows():
    n = 80
    data = pa.table({"k": [f"строка {i}" for i in range(n)], "color": [f"{i:06X}" for i in range(n)]})
    params = {"dataset": "t", "columns": [{"column": "k", "fill": "color"}], "min_font_size": 9}
    shape, ctx = fill_table(params, data)
    shown = len(shape.table.rows) - 1
    assert shown < n and any("показаны последние" in w for w in ctx.warnings)
    assert shape.table.cell(1, 0).text == f"строка {n - shown}"
    assert fills(shape, 0) == [f"{i:06X}" for i in range(n - shown, n)]


def test_table_fill_inherited_font_size():
    from autogenerator.contracts.ooxml import P

    def sizes(prepare) -> set[str]:
        shape, _ = fill_table({"dataset": "t"}, pa.table({"k": ["а"], "v": [1.0]}), prepare)
        return {r.get("sz") for r in shape.table._tbl.iter(f"{A}rPr")}

    def master_style(slide):
        master = slide.slide_layout.slide_master._element
        return master.find(f"{P}txStyles/{P}otherStyle/{A}lvl1pPr/{A}defRPr")

    def default_style(slide):
        pres = slide.part.package.presentation_part._element
        return pres.find(f"{P}defaultTextStyle/{A}lvl1pPr/{A}defRPr")

    def master_14(slide, shape):
        master_style(slide).set("sz", "1400")

    def presentation_16(slide, shape):
        master_style(slide).attrib.pop("sz")
        default_style(slide).set("sz", "1600")

    def nothing(slide, shape):
        master_style(slide).attrib.pop("sz")
        default_style(slide).attrib.pop("sz")

    assert sizes(master_14) == {"1400"}
    assert sizes(presentation_16) == {"1600"}
    assert sizes(nothing) == {"1800"}


def test_table_fill_keeps_template_row_heights():
    template_h = 438912  # высота строк таблицы синтетического шаблона

    def taller(slide, shape):
        trs = list(shape.table._tbl.tr_lst)
        trs[1].set("h", "500000")  # строка, с которой копируются новые
        trs[2].set("h", "600000")

    data = pa.table({"k": [f"строка {i}" for i in range(6)]})
    shape, _ = fill_table({"dataset": "t"}, data, taller)
    heights = [r.height for r in shape.table.rows]
    assert heights == [template_h, 500000, 600000, template_h, template_h, 500000, 500000]
    assert shape.height == sum(heights)

    # Высоты шаблона не помещаются до bottom, высоты по тексту — помещаются: кегль не уменьшается.
    four = pa.table({"k": [f"строка {i}" for i in range(4)]})
    top = Inches(1.5)
    bottom = (top + 5 * template_h - 100_000) / 914400
    shape, ctx = fill_table({"dataset": "t", "bottom": bottom}, four)
    assert all(r.height < template_h for r in shape.table.rows)
    assert {r.get("sz") for r in shape.table._tbl.iter(f"{A}rPr")} == {"1800"}
    assert not ctx.warnings


def test_series_labels_get_scenario_format():
    from lxml import etree

    from autogenerator.blocks_std.chart_xml import set_label_formats

    ser = (
        '<c:ser><c:idx val="{i}"/><c:order val="{i}"/><c:dLbls>{dlbl}<c:numFmt formatCode="#,##0" sourceLinked="0"/>'
        '<c:showVal val="1"/></c:dLbls></c:ser>'
    )
    dlbl = '<c:dLbl><c:idx val="0"/><c:numFmt formatCode="#,##0" sourceLinked="0"/><c:showVal val="1"/></c:dLbl>'
    cs = etree.fromstring(
        f'<c:chartSpace xmlns:c="{C[1:-1]}"><c:chart><c:plotArea><c:barChart>'
        + ser.format(i=0, dlbl=dlbl)
        + ser.format(i=1, dlbl="")
        + "</c:barChart></c:plotArea></c:chart></c:chartSpace>"
    )
    set_label_formats(cs, ["#,##0;-#,##0;", None])
    codes = [nf.get("formatCode") for nf in cs.iter(f"{C}numFmt")]
    assert codes == ["#,##0;-#,##0;", "#,##0;-#,##0;", "#,##0"]  # подпись точки, подписи серии, вторая серия


def test_point_unique_ids_shared_with_labels_but_not_between_points():
    from lxml import etree

    from autogenerator.blocks_std.chart_xml import C16, check_chart, cleanup_points

    def ext(uid: str) -> str:
        return f'<c:extLst><c:ext uri="{{C3BC}}"><c16:uniqueId val="{uid}"/></c:ext></c:extLst>'

    # Так круговую с настройками секторов сохраняет PowerPoint: у сектора и его подписи один id.
    dpt = '<c:dPt><c:idx val="{i}"/><c:spPr><a:solidFill><a:srgbClr val="{rgb}"/></a:solidFill></c:spPr>{ext}</c:dPt>'
    dlbl = '<c:dLbl><c:idx val="{i}"/><c:showVal val="1"/>{ext}</c:dLbl>'
    points = [("{00000001-AA}", "FF0000"), ("{00000003-AA}", "00FF00")]
    cs = etree.fromstring(
        f'<c:chartSpace xmlns:c="{C[1:-1]}" xmlns:a="{A[1:-1]}" xmlns:c16="{C16}"><c:chart><c:plotArea><c:pieChart>'
        '<c:ser><c:idx val="0"/><c:order val="0"/>'
        + "".join(dpt.format(i=i, rgb=rgb, ext=ext(u)) for i, (u, rgb) in enumerate(points))
        + "<c:dLbls>"
        + "".join(dlbl.format(i=i, ext=ext(u)) for i, (u, _) in enumerate(points))
        + "</c:dLbls>"
        + '<c:val><c:numRef><c:numCache><c:ptCount val="2"/></c:numCache></c:numRef></c:val>'
        + ext("{0000000C-AA}")
        + "</c:ser></c:pieChart></c:plotArea></c:chart></c:chartSpace>"
    )
    assert check_chart(cs, 2) == []
    cleanup_points(cs, 4, {})  # секторов стало больше: новые копируют оформление, но не id
    uids = [u.get("val") for d in cs.iter(f"{C}dPt") for u in d.iter(f"{{{C16}}}uniqueId")]
    assert len(uids) == 4 and len(set(uids)) == 4
    for el in cs.iter(f"{C}ptCount"):
        el.set("val", "4")
    assert check_chart(cs, 4) == []
