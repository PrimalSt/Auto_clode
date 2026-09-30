from dataclasses import dataclass, field
from datetime import date

import pyarrow as pa
import pytest
from pptx import Presentation

from autogenerator.blocks_std.chart import ChartBlock
from autogenerator.blocks_std.formats import NBSP
from autogenerator.blocks_std.table import TableBlock
from autogenerator.blocks_std.text import TextBlock, referenced_metrics, render_text
from autogenerator.contracts import BlockData, BlockTarget, Geometry, Period


@dataclass
class Ctx:
    period: Period = field(default_factory=lambda: Period.parse("2026-03"))
    scenario_name: str = "Тест"
    warnings: list[str] = field(default_factory=list)

    def warn(self, message: str) -> None:
        self.warnings.append(message)


def target(layout: int = 6):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    return prs, BlockTarget(slide=slide, geometry=Geometry(x=0, y=0, cx=6_000_000, cy=4_000_000))


def test_text_renders_metrics_and_period():
    text = "Выручка за {{ period.month_prep }}: {{ metrics.rev | money(scale='million', decimals=1) }}"
    assert render_text(text, {"rev": 26_229_878}, Ctx()) == f"Выручка за марте: 26,2{NBSP}млн{NBSP}₽"
    assert render_text("{{ metrics.rev }}", {"rev": 1234.5}, Ctx()) == f"1{NBSP}234,5"


def test_empty_metric_prints_no_data():
    assert render_text("{{ metrics.plan | percent }}", {"plan": None}, Ctx()) == "нет данных"
    assert render_text("{{ change_pct(metrics.a, metrics.b) | percent }}", {"a": 1, "b": 0}, Ctx()) == "нет данных"


def test_text_errors():
    with pytest.raises(ValueError, match="неизвестная переменная"):
        render_text("{{ metrics.nope }}", {}, Ctx())
    with pytest.raises(ValueError, match="ошибка в тексте"):
        render_text("{{ metrics.a ", {}, Ctx())


def test_referenced_metrics():
    assert referenced_metrics("{{ metrics.a | number }} {{ metrics['b'] }} {{ period.label }}") == {
        "a",
        "b",
    }
    assert TextBlock().data_needs(TextBlock().parse_params({"text": "{{ metrics.x }}"})).metrics == {"x"}


def test_text_block_into_placeholder_and_textbox():
    _, t = target(layout=1)
    t.placeholder = t.slide.placeholders[1]
    block = TextBlock()
    block.render(t, block.parse_params({"text": "строка 1\nстрока 2"}), BlockData(), Ctx())
    assert [p.text for p in t.placeholder.text_frame.paragraphs] == ["строка 1", "строка 2"]
    _, t2 = target()
    block.render(t2, block.parse_params({"text": "в рамке", "font_size": 20}), BlockData(), Ctx())
    assert t2.slide.shapes[-1].text_frame.text == "в рамке"


def test_chart_block():
    table = pa.table({"month": [date(2026, 1, 1), date(2026, 2, 1)], "revenue": [10.0, None]})
    _, t = target()
    block = ChartBlock()
    params = block.parse_params({"chart": "bar", "dataset": "d", "x": "month", "series": ["revenue"]})
    block.render(t, params, BlockData(datasets={"d": table}), Ctx())
    chart = t.slide.shapes[-1].chart
    assert list(chart.plots[0].categories) == ["янв. 2026", "февр. 2026"]
    assert list(chart.series[0].values) == [10.0, None]
    assert chart.value_axis.minimum_scale == 0
    xml = chart._chartSpace.xml
    assert '<c:orientation val="maxMin"/>' in xml and '<c:crosses val="max"/>' in xml


def test_chart_block_errors():
    block = ChartBlock()
    params = block.parse_params({"dataset": "d", "x": "nope", "series": ["revenue"]})
    _, t = target()
    with pytest.raises(ValueError, match="нет столбца «nope»"):
        block.render(t, params, BlockData(datasets={"d": pa.table({"revenue": [1]})}), Ctx())
    with pytest.raises(ValueError, match="ровно одна серия"):
        block.parse_params({"chart": "pie", "dataset": "d", "x": "a", "series": ["b", "c"]})


def test_table_block():
    table = pa.table(
        {
            "month": [date(2026, 1, 1), date(2026, 2, 1)],
            "revenue": [22_852_947.4, None],
            "share": [0.5, 0.25],
        }
    )
    _, t = target()
    block = TableBlock()
    params = block.parse_params(
        {
            "dataset": "d",
            "columns": [
                {"column": "month", "header": "Месяц"},
                {"column": "revenue", "header": "Выручка"},
                {"column": "share", "header": "Доля", "format": {"kind": "percent", "decimals": 0}},
            ],
        }
    )
    block.render(t, params, BlockData(datasets={"d": table}), Ctx())
    tbl = t.slide.shapes[-1].table
    rows = [[c.text for c in r.cells] for r in tbl.rows]
    assert rows == [
        ["Месяц", "Выручка", "Доля"],
        ["январь 2026", f"22{NBSP}852{NBSP}947", "50%"],
        ["февраль 2026", "—", "25%"],
    ]
