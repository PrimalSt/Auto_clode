"""render проверяется без theme и blocks-std: манифест шаблона собирается здесь по
стандартному шаблону python-pptx, а блоки — маленькие тестовые плагины."""

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pydantic import BaseModel

from autogenerator.contracts import (
    BlockPlugin,
    DataNeeds,
    Geometry,
    LayoutInfo,
    RoleBinding,
    SlotInfo,
    ThemeManifest,
)
from autogenerator.plugin_host import PluginRegistry

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _geometry(shape: Any) -> Geometry:
    return Geometry(x=shape.left, y=shape.top, cx=shape.width, cy=shape.height)


@pytest.fixture
def theme(tmp_path: Path) -> ThemeManifest:
    prs = Presentation()
    # Слайд шаблона: в отчёт он попасть не должен.
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Старый слайд шаблона"
    path = tmp_path / "template.pptx"
    prs.save(str(path))
    ids = [e.get("id") for e in prs.slide_master._element.find(f"{P}sldLayoutIdLst")]
    tac = prs.slide_layouts[1]
    title, body = tac.placeholders[0], tac.placeholders[1]
    return ThemeManifest(
        source_path=str(path),
        pptx_path=str(path),
        sha256="",
        slide_width=prs.slide_width,
        slide_height=prs.slide_height,
        layouts=[LayoutInfo(key=ids[1], name=tac.name, master=0)],
        roles=[
            RoleBinding(
                role="title_and_content",
                layout_key=ids[1],
                layout_name=tac.name,
                slots=[
                    SlotInfo(name="title", placeholder_idx=0, geometry=_geometry(title)),
                    SlotInfo(name="body", placeholder_idx=1, geometry=_geometry(body)),
                ],
            )
        ],
    )


class TextParams(BaseModel):
    text: str
    metric: str | None = None
    dataset: str | None = None


class Echo(BlockPlugin):
    """Пишет текст в плейсхолдер или в новую рамку; ``{metric}`` заменяется значением."""

    name = "echo"
    Params = TextParams

    def data_needs(self, params: Any) -> DataNeeds:
        return DataNeeds(
            metrics={params.metric} if params.metric else set(),
            datasets={params.dataset} if params.dataset else set(),
        )

    def render(self, target: Any, params: Any, data: Any, ctx: Any) -> None:
        text = params.text.replace("{metric}", str(data.metrics.get(params.metric)))
        if target.placeholder is not None:
            target.placeholder.text_frame.text = text
        else:
            g = target.geometry
            target.slide.shapes.add_textbox(g.x, g.y, g.cx, g.cy).text_frame.text = text


class Boom(BlockPlugin):
    name = "boom"

    def data_needs(self, params: Any) -> DataNeeds:
        return DataNeeds()

    def render(self, target: Any, params: Any, data: Any, ctx: Any) -> None:
        raise RuntimeError("блок сломался")


@pytest.fixture
def registry() -> PluginRegistry:
    return PluginRegistry.from_plugins([Echo, Boom])
