"""Слайды-образцы: сборка на синтетическом шаблоне examples/templates/synthetic.pptx.

В отличие от остальных тестов render, здесь работают настоящие theme (манифест шаблона) и
blocks-std (метки, chart_fill, table_fill): проверяется сборка целиком — копии, порядок,
удаление слайдов шаблона, разделы, имена частей и проверка после сборки."""

import io
import zipfile
from pathlib import Path

import pyarrow as pa
import pytest
from pptx import Presentation

from autogenerator.blocks_std.chart_fill import ChartFillBlock
from autogenerator.blocks_std.markers import MarkersBlock
from autogenerator.blocks_std.table_fill import TableFillBlock
from autogenerator.contracts import EngineResult, IssueLevel, Period, ScenarioSpec
from autogenerator.plugin_host import PluginRegistry
from autogenerator.render import build_presentation, validate_slides
from autogenerator.render.package_ops import section_names
from autogenerator.theme import import_template

TEMPLATE = Path(__file__).parents[3] / "examples" / "templates" / "synthetic.pptx"
MARCH = Period.parse("2026-03")
COVER, SUMMARY, REGIONS, PLAN, PAIR, TABLE, DRAFT, FINAL = range(256, 264)


@pytest.fixture(scope="module")
def theme(tmp_path_factory):
    return import_template(TEMPLATE, tmp_path_factory.mktemp("theme"))


@pytest.fixture(scope="module")
def registry():
    return PluginRegistry.from_plugins([MarkersBlock, ChartFillBlock, TableFillBlock])


def scenario(*slides, markers=None) -> ScenarioSpec:
    return ScenarioSpec.model_validate(
        {
            "name": "Тест",
            "inputs": [{"id": "sales", "source": "s"}],
            "datasets": [{"id": d, "input": "sales"} for d in ("months", "regions", "table")],
            "metrics": [{"id": m, "input": "sales", "fn": "sum", "column": "x"} for m in ("rev", "growth", "share")],
            "markers": markers if markers is not None else {"Месяц": "period.month", "Год": "period.year"},
            "slides": list(slides),
        }
    )


SUMMARY_SLIDE = {
    "example": SUMMARY,
    "markers": {
        "Выручка": {"metric": "rev", "scale": "million", "decimals": 1},
        "Прирост+": {"metric": "growth", "percent": True, "sign": True},
        "Доля": {"metric": "share", "percent": True, "decimals": 0},
    },
    "blocks": [
        {
            "type": "chart_fill",
            "shape": 6,
            "dataset": "months",
            "categories": "m",
            "series": ["a", "b", {"column": "s", "number_format": "0%"}],
        }
    ],
}
REGIONS_SLIDE = {
    "example": REGIONS,
    "markers": {"Доля": {"metric": "share", "percent": True, "decimals": 0}},
    "blocks": [{"type": "chart_fill", "shape": 4, "dataset": "regions", "categories": "r", "series": ["v"]}],
}
TABLE_SLIDE = {"example": TABLE, "blocks": [{"type": "table_fill", "shape": 3, "dataset": "table"}]}


def engine(**metrics) -> EngineResult:
    return EngineResult(
        period=MARCH,
        metrics={"rev": 26_229_878, "growth": 0.083, "share": 0.35, **metrics},
        datasets={
            "months": pa.table({"m": ["янв", "фев", "мар"], "a": [1.0, 2, 3], "b": [4.0, 5, 6], "s": [0.2, 0.3, 0.4]}),
            "regions": pa.table({"r": ["Север", "Юг", "Запад"], "v": [5.0, 3, 2]}),
            "table": pa.table({"Регион": ["Север", "Юг"], "2026-01": [1.5, 2.5], "2026-02": [1.7, 2.1]}),
        },
    )


def texts(slide) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_examples_are_filled_ordered_and_unused_removed(theme, registry, tmp_path: Path):
    sc = scenario(TABLE_SLIDE, SUMMARY_SLIDE, {"example": COVER}, REGIONS_SLIDE, {"example": FINAL})
    assert [i for i in validate_slides(sc, registry, theme) if i.level == IssueLevel.ERROR] == []
    res = build_presentation(sc, theme, engine(), registry, tmp_path / "out.pptx")
    assert res.output_path, res.issues
    prs = Presentation(res.output_path)
    titles = [texts(s).splitlines()[0] for s in prs.slides]
    assert titles == ["Выручка регионов по месяцам", "Итоги, март 2026", "Продажи", "Регионы, март 2026", "Спасибо!"]
    body = texts(prs.slides[1])
    assert "26,2 млн ₽" in body and "+8,3%" in body and "35% плана" in body
    table = next(s for s in prs.slides[0].shapes if s.has_table).table
    assert [c.text for c in table.rows[0].cells] == ["Регион", "янв. 2026", "февр. 2026"]
    chart = next(s for s in prs.slides[3].shapes if s.has_chart).chart
    assert list(chart.plots[0].categories) == ["Север", "Юг", "Запад"]
    # Части слайдов — по порядку, разделы — по итоговым слайдам, авторов примечаний нет.
    names = zipfile.ZipFile(res.output_path).namelist()
    assert sorted(n for n in names if n.startswith("ppt/slides/slide")) == [
        f"ppt/slides/slide{i}.xml" for i in range(1, 6)
    ]
    assert "ppt/commentAuthors.xml" not in names
    ids = [s.slide_id for s in prs.slides]
    assert set(section_names(prs)) == set(ids)
    assert section_names(prs)[ids[-1]] == "Служебное"


def test_example_used_twice_gets_own_chart_and_workbook(theme, registry, tmp_path: Path):
    second = {**REGIONS_SLIDE, "blocks": [{**REGIONS_SLIDE["blocks"][0], "dataset": "regions2"}]}
    sc = scenario(REGIONS_SLIDE, second)
    sc.datasets.append(sc.datasets[0].model_copy(update={"id": "regions2"}))
    eng = engine()
    eng.datasets["regions2"] = pa.table({"r": ["А", "Б"], "v": [1.0, 2]})
    res = build_presentation(sc, theme, eng, registry, tmp_path / "out.pptx")
    assert res.output_path, res.issues
    prs = Presentation(res.output_path)
    charts = [next(s for s in slide.shapes if s.has_chart).chart for slide in prs.slides]
    assert [list(c.plots[0].categories) for c in charts] == [["Север", "Юг", "Запад"], ["А", "Б"]]
    parts = {c.part.partname for c in charts}
    books = {c.part.chart_workbook.xlsx_part.partname for c in charts}
    assert len(parts) == 2 and len(books) == 2
    # Книга первого графика не перезаписана данными копии.
    book = zipfile.ZipFile(io.BytesIO(charts[0].part.chart_workbook.xlsx_part.blob))
    strings = book.read("xl/sharedStrings.xml").decode()
    assert "Запад" in strings and ">Б<" not in strings
    assert len(set(section_names(prs).values())) == 1


def test_unbound_marker_stops_build_but_not_preview(theme, registry, tmp_path: Path):
    sc = scenario({"example": COVER}, markers={"Месяц": "period.month"})
    issues = validate_slides(sc, registry, theme)
    assert any("{{Год}} не привязана" in i.message and i.level == IssueLevel.ERROR for i in issues)
    res = build_presentation(sc, theme, engine(), registry, tmp_path / "out.pptx")
    assert res.output_path is None
    res = build_presentation(sc, theme, engine(), registry, tmp_path / "preview.pptx", preview=True)
    assert res.output_path
    assert "март {{Год}}" in texts(Presentation(res.output_path).slides[0])


def test_empty_metric_stops_with_all_places(theme, registry, tmp_path: Path):
    sc = scenario(SUMMARY_SLIDE, REGIONS_SLIDE)
    res = build_presentation(sc, theme, engine(share=None), registry, tmp_path / "out.pptx")
    assert res.output_path is None
    errors = [i for i in res.issues if i.level == IssueLevel.ERROR]
    assert {i.node for i in errors} == {"slide:1", "slide:2"}
    assert all("{{Доля}}" in i.message for i in errors)


def test_kept_marker_passes_leftover_check(theme, registry, tmp_path: Path):
    sc = scenario({"example": COVER, "markers": {"Год": {"metric": "rev", "empty": "keep"}}})
    res = build_presentation(sc, theme, engine(rev=None), registry, tmp_path / "out.pptx")
    assert res.output_path, res.issues
    assert "март {{Год}}" in texts(Presentation(res.output_path).slides[0])


def test_unfilled_chart_needs_block_or_keep(theme, registry, tmp_path: Path):
    bare = {k: v for k, v in REGIONS_SLIDE.items() if k != "blocks"}
    messages = [i.message for i in validate_slides(scenario(bare), registry, theme)]
    assert any("не заполнен" in m and "keep" in m for m in messages)
    res = build_presentation(scenario({**bare, "keep": [4]}), theme, engine(), registry, tmp_path / "out.pptx")
    assert res.output_path, res.issues
    chart = next(s for s in Presentation(res.output_path).slides[0].shapes if s.has_chart).chart
    assert len(list(chart.plots[0].categories)) == 6  # демо-данные шаблона остались намеренно


def test_wrong_blocks_on_example(theme, registry):
    sc = scenario(
        {"example": 999},
        {
            "example": SUMMARY,
            "blocks": [{"type": "chart_fill", "dataset": "months", "categories": "m", "series": ["a"]}],
        },
        {"layout": "title_and_content", "blocks": [{"type": "table_fill", "dataset": "table"}]},
    )
    messages = [i.message for i in validate_slides(sc, registry, theme)]
    assert any("нет слайда с id 999" in m for m in messages)
    assert any("не указана фигура" in m for m in messages)
    assert any("заполняет фигуру слайда-образца" in m for m in messages)


def test_only_one_slide_for_preview(theme, registry, tmp_path: Path):
    sc = scenario({"example": COVER}, SUMMARY_SLIDE, {"example": FINAL})
    res = build_presentation(sc, theme, engine(), registry, tmp_path / "one.pptx", preview=True, only=2)
    assert res.output_path and res.slides == 1
    assert [n.id for n in res.nodes] == ["slide:2"]
