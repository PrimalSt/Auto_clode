from pathlib import Path

import pytest
from pptx import Presentation

from autogenerator.contracts import (
    EngineResult,
    IssueLevel,
    NodeKind,
    NodeState,
    NodeStatus,
    Period,
    ScenarioSpec,
)
from autogenerator.render import build as build_module
from autogenerator.render import build_presentation, sanitize_filename, validate_slides

MARCH = Period.parse("2026-03")


def scenario(*slides, metrics=("rev",)) -> ScenarioSpec:
    return ScenarioSpec.model_validate(
        {
            "name": "Тест",
            "inputs": [{"id": "sales", "source": "s"}],
            "metrics": [{"id": m, "input": "sales", "fn": "sum", "column": "x"} for m in metrics],
            "slides": list(slides),
        }
    )


def slide(*blocks) -> dict:
    return {"layout": "title_and_content", "blocks": list(blocks)}


def texts(path: str) -> list[list[str]]:
    prs = Presentation(path)
    return [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in prs.slides]


def test_builds_slides_from_layouts(theme, registry, tmp_path: Path):
    sc = scenario(
        slide(
            {"type": "echo", "slot": "title", "text": "Итоги"},
            {"type": "echo", "text": "Выручка {metric}", "metric": "rev"},
        ),
        slide({"type": "echo", "slot": "title", "text": "Только заголовок"}),
    )
    res = build_presentation(
        sc, theme, EngineResult(period=MARCH, metrics={"rev": 42}), registry, tmp_path / "out.pptx"
    )
    assert res.output_path and res.slides == 2
    # Слайд шаблона удалён, пустой плейсхолдер второго слайда — тоже.
    assert texts(res.output_path) == [["Итоги", "Выручка 42"], ["Только заголовок"]]
    assert Presentation(res.output_path).core_properties.title == "Тест: 2026-03"


def test_block_error_marks_only_its_slide(theme, registry, tmp_path: Path):
    sc = scenario(slide({"type": "boom"}), slide({"type": "echo", "text": "цел"}))
    res = build_presentation(sc, theme, EngineResult(period=MARCH), registry, tmp_path / "out.pptx")
    assert res.output_path
    assert [n.state for n in res.nodes] == [NodeState.ERROR, NodeState.OK]
    assert texts(res.output_path) == [["Ошибка: блок сломался"], ["цел"]]


def test_block_with_failed_dependency_gets_marker(theme, registry, tmp_path: Path):
    failed = NodeStatus(id="metric:rev", kind=NodeKind.METRIC, state=NodeState.SKIPPED, blocked_by="input:sales")
    engine = EngineResult(period=MARCH, nodes=[failed])
    sc = scenario(slide({"type": "echo", "text": "{metric}", "metric": "rev"}))
    res = build_presentation(sc, theme, engine, registry, tmp_path / "out.pptx")
    assert texts(res.output_path) == [["Ошибка: зависит от «input:sales», в котором ошибка"]]


def test_block_without_its_data_gets_marker(theme, registry, tmp_path: Path):
    sc = scenario(slide({"type": "echo", "text": "{metric}", "metric": "rev"}))
    res = build_presentation(sc, theme, EngineResult(period=MARCH), registry, tmp_path / "out.pptx")
    assert texts(res.output_path) == [["Ошибка: в итогах расчёта нет показателя «rev»"]]


def test_leftover_marker_is_not_saved(theme, registry, tmp_path: Path):
    sc = scenario(slide({"type": "echo", "text": "Выручка {{ metrics.rev }}"}))
    res = build_presentation(sc, theme, EngineResult(period=MARCH), registry, tmp_path / "out.pptx")
    assert res.output_path is None
    assert any("{{" in i.message for i in res.issues)
    assert not (tmp_path / "out.pptx").exists()


def test_validation(theme, registry):
    sc = scenario(
        {"layout": "section", "blocks": []},
        slide(
            {"type": "nope"},
            {"type": "echo", "text": "x", "metric": "missing"},
            {"type": "echo", "slot": "left", "text": "x"},
        ),
        {"example": {"id": 256}},
        slide({"type": "echo"}),
    )
    messages = [i.message for i in validate_slides(sc, registry, theme)]
    assert any("нет макета для роли «section»" in m for m in messages)
    assert any("Нет плагина block «nope»" in m for m in messages)
    assert any("нет показателя «missing»" in m for m in messages)
    assert any("нет области «left»" in m for m in messages)
    assert any("нет слайда с id 256" in m for m in messages)
    assert any("параметры блока «echo»" in m and "text" in m for m in messages)


def test_busy_file_is_saved_next_to_it(theme, registry, tmp_path: Path, monkeypatch):
    target = tmp_path / "Отчёт.pptx"
    target.write_bytes(b"open in PowerPoint")
    monkeypatch.setattr(build_module, "_busy", lambda p: p == target)
    res = build_presentation(
        scenario(slide({"type": "echo", "text": "x"})),
        theme,
        EngineResult(period=MARCH),
        registry,
        target,
    )
    assert res.output_path == str(tmp_path / "Отчёт (2).pptx")
    assert target.read_bytes() == b"open in PowerPoint"
    assert any(i.level == IssueLevel.WARNING and "открыт" in i.message for i in res.issues)
    assert not list(tmp_path.glob(".*.tmp"))


@pytest.mark.parametrize(
    ("name", "safe"),
    [
        ("Отчёт: март/2026?", "Отчёт_ март_2026_"),
        ("CON", "_CON"),
        ("отчёт. ", "отчёт"),
        ("", "Отчёт"),
    ],
)
def test_sanitize_filename(name, safe):
    assert sanitize_filename(name) == safe


def _png() -> bytes:
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    raw = zlib.compress(b"\x00\xff\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", raw)
        + chunk(b"IEND", b"")
    )


def test_design_elements_are_copied_to_new_slides(theme, registry, tmp_path: Path):
    import io

    from pptx.util import Inches

    from autogenerator.contracts import DesignElement, Geometry

    prs = Presentation(theme.pptx_path)
    cover = prs.slides[0]
    pic = cover.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(6), Inches(1), Inches(0.5))
    pic.name = "Логотип"
    sid = int(prs.slides._sldIdLst[0].get("id"))
    path = tmp_path / "with_logo.pptx"
    prs.save(str(path))
    where = Geometry(x=int(Inches(2)), y=int(Inches(6.5)), cx=int(Inches(1)), cy=int(Inches(0.5)))
    logo = DesignElement(name="Логотип", kind="picture", shape_id=pic.shape_id, from_slide=sid, geometry=where)
    role = theme.roles[0].model_copy(update={"decorations": [logo]})
    th = theme.model_copy(update={"pptx_path": str(path), "roles": [role]})
    sc = scenario(slide({"type": "echo", "slot": "title", "text": "Итоги"}), slide({"type": "echo", "text": "x"}))
    res = build_presentation(sc, th, EngineResult(period=MARCH, metrics={"rev": 1}), registry, tmp_path / "out.pptx")
    assert res.output_path and res.slides == 2
    out = Presentation(res.output_path)
    for s in out.slides:  # слайд шаблона удалён, а картинка осталась на новых слайдах
        pics = [sh for sh in s.shapes if sh.name == "Логотип"]
        assert len(pics) == 1 and (pics[0].left, pics[0].top) == (where.x, where.y)
        assert pics[0].image.blob == _png() and len({sh.shape_id for sh in s.shapes}) == len(s.shapes)
