"""Операции со слайдами и частями файла, которых нет в python-pptx (ARCHITECTURE.md, раздел 6.5).

- Новый слайд и копия слайда-образца получают свободное имя части: python-pptx называет
  часть по числу слайдов (``slideN.xml``) и не проверяет, занято ли имя.
- Копия слайда-образца копирует вместе с ним графики, их книги Excel, стили и цвета графиков,
  заметки, теги и примечания: иначе ``replace_data()`` на копии перезапишет книгу оригинала.
  Макет, картинки и внешние ссылки остаются общими.
- Удалённый слайд уходит вместе со связью; ссылки других слайдов на него снимаются, иначе
  часть удалённого слайда осталась бы в файле.
- Перед сохранением части слайдов и заметок переименовываются по порядку, список разделов
  (``p14:sectionLst``) и ``docProps/app.xml`` переписываются по итоговым слайдам, список
  авторов примечаний без единого примечания удаляется.
"""

from __future__ import annotations

import copy
import random
import re
import uuid
from typing import Any

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, PartFactory, _Relationship
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlidePart

from autogenerator.contracts.ooxml import NS, A, P, R, xp

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTIONS_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
EP = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
VT = "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}"
MS_CHART_STYLE = "http://schemas.microsoft.com/office/2011/relationships/chartStyle"
MS_CHART_COLORS = "http://schemas.microsoft.com/office/2011/relationships/chartColorStyle"
MS_DIAGRAM_DRAWING = "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing"
MS_COMMENTS = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
MS_AUTHORS = "http://schemas.microsoft.com/office/2018/10/relationships/authors"

COPIED = {
    RT.CHART,
    RT.PACKAGE,
    RT.OLE_OBJECT,
    RT.NOTES_SLIDE,
    RT.TAGS,
    RT.COMMENTS,
    RT.THEME_OVERRIDE,
    RT.CHART_USER_SHAPES,
    RT.DIAGRAM_DATA,
    RT.DIAGRAM_LAYOUT,
    RT.DIAGRAM_QUICK_STYLE,
    RT.DIAGRAM_COLORS,
    RT.VML_DRAWING,
    MS_CHART_STYLE,
    MS_CHART_COLORS,
    MS_DIAGRAM_DRAWING,
    MS_COMMENTS,
}
"""Связи, чьи части копируются вместе со слайдом. Остальные (макет, картинки, мастер
заметок, ссылки) копия делит с оригиналом."""


def _template(partname: str) -> str:
    """Шаблон имени для ``next_partname``: ``/ppt/charts/chart12.xml`` → ``/ppt/charts/chart%d.xml``."""
    name = partname.replace("%", "%%")
    m = re.match(r"^(.*?)(\d+)(\.\w+)$", name)
    if m:
        return f"{m.group(1)}%d{m.group(3)}"
    m = re.match(r"^(.*)(\.\w+)$", name)
    return f"{m.group(1)}%d{m.group(2)}" if m else f"{name}%d"


def _set_rel(part: Part, rId: str, reltype: str, target: Part | str, external: bool) -> None:
    """Связь с заданным rId: XML копии ссылается на те же rId, что и оригинал."""
    rels = part.rels
    rels._rels[rId] = _Relationship(rels._base_uri, rId, reltype, RTM.EXTERNAL if external else RTM.INTERNAL, target)


def _copy_part(src: Part, partname: PackURI, package: Any) -> Part:
    return PartFactory(partname, src.content_type, package, src.blob)


def _copy_rels(src: Part, dst: Part, copies: dict[Part, Part], package: Any) -> None:
    """Связи копии части: части из ``COPIED`` копируются (рекурсивно), остальные общие.
    Ссылка на уже скопированную часть (заметки → слайд) ведёт на её копию."""
    for rId, rel in src.rels.items():
        if rel.is_external:
            _set_rel(dst, rId, rel.reltype, rel.target_ref, external=True)
            continue
        target = rel.target_part
        if target in copies:
            _set_rel(dst, rId, rel.reltype, copies[target], external=False)
            continue
        if rel.reltype not in COPIED:
            _set_rel(dst, rId, rel.reltype, target, external=False)
            continue
        # Связь ставится сразу, до копирования детей: тогда новая часть видна в пакете
        # и next_partname не выдаст её имя ещё раз.
        new = _copy_part(target, package.next_partname(_template(str(target.partname))), package)
        copies[target] = new
        _set_rel(dst, rId, rel.reltype, new, external=False)
        _copy_rels(target, new, copies, package)


def _new_creation_id(slide_el: Any) -> None:
    for el in xp(slide_el, ".//p14:creationId"):
        el.set("val", str(random.randint(1, 2**32 - 1)))


def add_slide(prs: Any, layout: Any) -> Any:
    """Новый слайд на макете ``layout`` в конце презентации, со свободным именем части."""
    package = prs.part.package
    part = SlidePart.new(package.next_partname("/ppt/slides/slide%d.xml"), package, layout.part)
    rId = prs.part.relate_to(part, RT.SLIDE)
    prs.slides._sldIdLst.add_sldId(rId)
    slide = part.slide
    slide.shapes.clone_layout_placeholders(layout)
    return slide


def find_shape(tree: Any, shape_id: int) -> Any | None:
    """Фигура верхнего уровня дерева фигур по id."""
    for el in tree:
        c = el.find(f".//{P}cNvPr")
        if c is not None and c.get("id") == str(shape_id) and el.getparent() is tree:
            return el
    return None


def copy_shape(src_part: Part, el: Any, slide: Any, x: int, y: int) -> Any:
    """Копия фигуры (картинки, группы, фигуры без данных) со слайда или макета на ``slide``:
    связи с картинками и ссылками ставятся заново, id фигур — свободные, левый верхний угол —
    (x, y). Копия ложится под остальные фигуры: элемент оформления — фон для содержания."""
    new = copy.deepcopy(el)
    dst: Part = slide.part
    for d in new.iter():
        for attr in [a for a in d.attrib if a.startswith(R)]:
            rel = src_part.rels.get(d.get(attr))
            if rel is None:
                del d.attrib[attr]
            elif rel.is_external:
                d.set(attr, dst.relate_to(rel.target_ref, rel.reltype, is_external=True))
            else:
                d.set(attr, dst.relate_to(rel.target_part, rel.reltype))
    tree = slide.shapes._spTree
    ids = [int(c.get("id", 0)) for c in slide._element.iter(f"{P}cNvPr")]
    next_id = max(ids, default=1) + 1
    for c in new.iter(f"{P}cNvPr"):
        c.set("id", str(next_id))
        next_id += 1
    xfrm = new.find(f"{P}grpSpPr/{A}xfrm") if new.tag == f"{P}grpSp" else new.find(f"{P}spPr/{A}xfrm")
    off = xfrm.find(f"{A}off") if xfrm is not None else None
    if off is not None:
        off.set("x", str(x))
        off.set("y", str(y))
    tree.insert(2, new)  # после nvGrpSpPr и grpSpPr
    return new


def duplicate_slide(prs: Any, slide: Any) -> Any:
    """Копия слайда в конце презентации вместе с графиками, книгами, заметками и тегами."""
    package = prs.part.package
    src: Part = slide.part
    new = _copy_part(src, package.next_partname("/ppt/slides/slide%d.xml"), package)
    rId = prs.part.relate_to(new, RT.SLIDE)
    prs.slides._sldIdLst.add_sldId(rId)
    _copy_rels(src, new, {src: new}, package)
    assert isinstance(new, SlidePart)
    _new_creation_id(new._element)
    return new.slide


def slide_entries(prs: Any) -> list[tuple[int, Any]]:
    """(sldId, слайд) по порядку."""
    lst = prs.slides._sldIdLst
    return [(int(e.get("id")), prs.part.related_part(e.get(f"{R}id")).slide) for e in lst]


def remove_slides(prs: Any, slide_ids: set[int]) -> None:
    """Удалить слайды вместе со связями; ссылки оставшихся слайдов на удалённые снимаются."""
    lst = prs.slides._sldIdLst
    gone: set[Part] = set()
    for e in list(lst):
        if int(e.get("id")) in slide_ids:
            rId = e.get(f"{R}id")
            gone.add(prs.part.related_part(rId))
            lst.remove(e)
            prs.part.drop_rel(rId)
    if not gone:
        return
    for e in lst:
        part = prs.part.related_part(e.get(f"{R}id"))
        for rId, rel in list(part.rels.items()):
            if not rel.is_external and rel.target_part in gone:
                for el in xp(part._element, f".//*[@r:id='{rId}']"):
                    el.getparent().remove(el)
                part.rels._rels.pop(rId)


def reorder_slides(prs: Any, slide_ids: list[int]) -> None:
    lst = prs.slides._sldIdLst
    by_id = {int(e.get("id")): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    for sid in slide_ids:
        lst.append(by_id[sid])


def renumber_parts(prs: Any) -> None:
    """Части слайдов и заметок — ``slide1.xml``…, ``notesSlide1.xml``… по порядку слайдов."""
    notes = 0
    for i, (_, slide) in enumerate(slide_entries(prs), start=1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
        if slide.has_notes_slide:
            notes += 1
            slide.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{notes}.xml")


# --- разделы, app.xml, авторы примечаний ---------------------------------------------------


def _section_list(prs: Any) -> Any | None:
    found = xp(prs.part._element, f"p:extLst/p:ext[@uri='{SECTIONS_URI}']/p14:sectionLst")
    return found[0] if found else None


def section_names(prs: Any) -> dict[int, str]:
    """Раздел каждого слайда: sldId → имя раздела."""
    lst = _section_list(prs)
    out: dict[int, str] = {}
    if lst is None:
        return out
    for sec in lst:
        for sid in xp(sec, "p14:sldIdLst/p14:sldId"):
            out[int(sid.get("id"))] = sec.get("name", "")
    return out


def rewrite_sections(prs: Any, origin: dict[int, int | None]) -> None:
    """Переписать разделы по итоговому списку слайдов. Слайд-образец и его копии остаются в
    разделе оригинала, новый слайд — в разделе предыдущего слайда. Разделы идут подряд: если
    слайды разных разделов чередуются, раздел повторяется с тем же именем; пустые удаляются."""
    lst = _section_list(prs)
    if lst is None:
        return
    sec_of: dict[int, Any] = {}
    for sec in lst:
        for sid in xp(sec, "p14:sldIdLst/p14:sldId"):
            sec_of[int(sid.get("id"))] = sec
    sections = list(lst)
    final = [sid for sid, _ in slide_entries(prs)]
    groups: list[tuple[Any, list[int]]] = []
    prev = sections[0] if sections else None
    for sid in final:
        src = origin.get(sid)
        sec = sec_of.get(src) if src is not None else None
        sec = sec if sec is not None else prev
        if groups and groups[-1][0] is sec:
            groups[-1][1].append(sid)
        else:
            groups.append((sec, [sid]))
        prev = sec
    for sec in sections:
        lst.remove(sec)
    if not groups:
        ext = lst.getparent()
        ext.getparent().remove(ext)
        return
    used: set[str] = set()
    for sec, ids in groups:
        new = etree.SubElement(lst, f"{{{P14}}}section")
        new.set("name", sec.get("name", "") if sec is not None else "")
        sec_id = sec.get("id") if sec is not None else None
        if sec_id is None or sec_id in used:
            sec_id = "{" + str(uuid.uuid4()).upper() + "}"
        used.add(sec_id)
        new.set("id", sec_id)
        ids_el = etree.SubElement(new, f"{{{P14}}}sldIdLst")
        for sid in ids:
            etree.SubElement(ids_el, f"{{{P14}}}sldId", id=str(sid))


def _package_part(prs: Any, reltype: str) -> Part | None:
    for rel in prs.part.package._rels.values():
        if rel.reltype == reltype and not rel.is_external:
            return rel.target_part
    return None


def rewrite_app_xml(prs: Any, titles: list[str], template_slides: int) -> None:
    """``docProps/app.xml``: число слайдов, заметок и скрытых слайдов и названия слайдов.
    Названия — последняя группа в ``TitlesOfParts``, если в ней было ровно столько записей,
    сколько слайдов в шаблоне (так пишет PowerPoint); иначе список названий удаляется."""
    part = _package_part(prs, RT.EXTENDED_PROPERTIES)
    if part is None:
        return
    root = etree.fromstring(part.blob)
    slides = slide_entries(prs)

    def put(tag: str, value: int) -> None:
        el = root.find(f"{EP}{tag}")
        if el is not None:
            el.text = str(value)

    put("Slides", len(slides))
    put("Notes", sum(1 for _, s in slides if s.has_notes_slide))
    put("HiddenSlides", sum(1 for _, s in slides if s._element.get("show") == "0"))
    heads = root.find(f"{EP}HeadingPairs/{VT}vector")
    parts = root.find(f"{EP}TitlesOfParts/{VT}vector")
    ok = heads is not None and parts is not None and len(heads) >= 2
    if ok:
        count_el = heads[-1].find(f"{VT}i4")
        ok = count_el is not None and int(count_el.text or 0) == template_slides and len(parts) >= template_slides
    if not ok:
        for tag in ("HeadingPairs", "TitlesOfParts"):
            el = root.find(f"{EP}{tag}")
            if el is not None:
                root.remove(el)
    else:
        assert heads is not None and parts is not None and count_el is not None
        for el in list(parts)[len(parts) - template_slides :]:
            parts.remove(el)
        for t in titles:
            etree.SubElement(parts, f"{VT}lpstr").text = t
        parts.set("size", str(len(parts)))
        count_el.text = str(len(titles))
        if not titles:  # пустая группа не пишется
            heads.remove(heads[-1])
            heads.remove(heads[-1])
            heads.set("size", str(len(heads)))
    part.blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def drop_unused_comment_authors(prs: Any) -> None:
    """Список авторов примечаний, на который не ссылается ни одно примечание, удаляется."""
    has_old = has_new = False
    for _, slide in slide_entries(prs):
        for rel in slide.part.rels.values():
            has_old |= rel.reltype == RT.COMMENTS
            has_new |= rel.reltype == MS_COMMENTS
    for rId, rel in list(prs.part.rels.items()):
        if (rel.reltype == RT.COMMENT_AUTHORS and not has_old) or (rel.reltype == MS_AUTHORS and not has_new):
            prs.part.rels._rels.pop(rId)


# --- проверка пакета ----------------------------------------------------------------------


def check_package(prs: Any) -> list[str]:
    """Повторяющиеся имена частей, ссылки XML на несуществующие связи, слайды вне разделов."""
    problems: list[str] = []
    names: dict[str, int] = {}
    for part in prs.part.package.iter_parts():
        names[str(part.partname)] = names.get(str(part.partname), 0) + 1
        el = getattr(part, "_element", None)
        if el is None:
            continue
        for attr in xp(el, ".//@r:*"):
            if attr.attrname.startswith(R) and str(attr) and str(attr) not in part.rels:
                problems.append(f"{part.partname}: ссылка {attr} без связи")
    problems += [f"часть {n} повторяется" for n, k in names.items() if k > 1]
    lst = _section_list(prs)
    if lst is not None:
        in_sections = {int(s.get("id")) for s in xp(lst, ".//p14:sldId")}
        missing = [sid for sid, _ in slide_entries(prs) if sid not in in_sections]
        if missing:
            problems.append(f"слайды {', '.join(map(str, missing))} не входят ни в один раздел")
    return problems


__all__ = [
    "NS",
    "P",
    "add_slide",
    "check_package",
    "drop_unused_comment_authors",
    "duplicate_slide",
    "remove_slides",
    "renumber_parts",
    "reorder_slides",
    "rewrite_app_xml",
    "rewrite_sections",
    "section_names",
    "slide_entries",
]
