"""Сборка презентации в рабочей копии шаблона (ARCHITECTURE.md, раздел 6.5).

Отчёт никогда не собирается из пустого ``Presentation()``: всегда открывается копия шаблона,
поэтому сохраняются мастера, макеты, тема и шрифты. Слайд отчёта — новый слайд на макете
шаблона (блоки ставятся в области макета) или слайд-образец шаблона: метки в его тексте
заменяются значениями, графики и таблица получают данные, остальное остаётся как в шаблоне.
Неиспользуемые слайды шаблона удаляются, слайд-образец, взятый второй раз, копируется.

Ошибка блока на слайде из макета не валит отчёт: на слайд ставится пометка «Ошибка: …»,
запись идёт в журнал. На слайде-образце незаполненная метка, график или таблица
останавливает сборку: иначе в отчёте останутся демо-данные шаблона. Остальные слайды при
этом собираются, и все такие места показываются одним списком.
"""

from __future__ import annotations

import os
import re
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.shapes.group import GroupShape
from pptx.util import Emu, Pt
from pydantic import ValidationError

from autogenerator.contracts import (
    AgenError,
    BlockData,
    BlockError,
    BlockPlugin,
    BlockTarget,
    BlockTargetKind,
    EngineResult,
    ErrorCode,
    Geometry,
    Issue,
    IssueLevel,
    NodeKind,
    NodeState,
    NodeStatus,
    Period,
    RenderResult,
    ScenarioSpec,
    SlideSpec,
    TemplateSlideInfo,
    ThemeManifest,
)
from autogenerator.contracts.ooxml import LOOSE_RE, P, all_text, iter_shapes, xp
from autogenerator.contracts.yaml_io import validation_message
from autogenerator.plugin_host import PluginRegistry

from .package_ops import (
    add_slide,
    check_package,
    drop_unused_comment_authors,
    duplicate_slide,
    remove_slides,
    renumber_parts,
    reorder_slides,
    rewrite_app_xml,
    rewrite_sections,
    slide_entries,
)

P_NS = P
DEFAULT_SLOT = "body"
MARKERS_BLOCK = "markers"
ERROR_COLOR = RGBColor(0xC0, 0x00, 0x00)
_KIND_NAMES = {BlockTargetKind.CHART: "график", BlockTargetKind.TABLE: "таблица"}


def _marker_name(key: Any) -> str:
    k = str(key).strip()
    if k.startswith("{{") and k.endswith("}}"):
        k = k[2:-2]
    return k.strip()


def _first_line(e: Exception) -> str:
    return str(e).splitlines()[0] if str(e) else type(e).__name__


# --- проверка без данных ---------------------------------------------------------------


def validate_slides(
    scenario: ScenarioSpec, registry: PluginRegistry, theme: ThemeManifest | None = None, *, preview: bool = False
) -> list[Issue]:
    """Проверить слайды: макеты и слайды-образцы, блоки и их параметры, области и фигуры,
    ссылки на наборы и показатели, привязку меток и незаполненные графики и таблицы.

    ``preview`` — для пробной сборки: непривязанные метки и незаполненные фигуры — предупреждения."""
    issues: list[Issue] = []
    datasets = {d.id for d in scenario.datasets}
    metrics = {m.id for m in scenario.metrics} | {c for m in scenario.metrics for c in m.compare_ids()}

    def err(node: str, msg: str, soft: bool = False) -> None:
        level = IssueLevel.WARNING if soft and preview else IssueLevel.ERROR
        issues.append(Issue(level=level, node=node, message=msg))

    def needs_ok(node: str, plugin: BlockPlugin, params: Any) -> None:
        needs = plugin.data_needs(params)
        for d in sorted(needs.datasets - datasets):
            err(node, f"нет набора «{d}»")
        for m in sorted(needs.metrics - metrics):
            err(node, f"нет показателя «{m}»")

    def parse(node: str, type_: str, raw: dict[str, Any], version: int = 1) -> tuple[BlockPlugin, Any] | None:
        try:
            plugin = registry.block(type_)
            return plugin, plugin.parse_params(raw, version)
        except AgenError as e:
            err(node, e.message)
        except ValidationError as e:
            err(node, validation_message(e, f"параметры блока «{type_}»"))
        except Exception as e:
            err(node, _first_line(e))
        return None

    has_examples = any(s.example is not None for s in scenario.slides if s.enabled)
    if scenario.markers or has_examples:
        parse("markers", MARKERS_BLOCK, {"common": scenario.markers})

    for n, slide in enumerate((s for s in scenario.slides if s.enabled), start=1):
        snode = f"slide:{n}"
        if slide.example is not None:
            _validate_example(scenario, slide, snode, theme, err, parse, needs_ok)
            continue
        binding = theme.role(slide.layout) if theme and slide.layout else None
        if theme is not None and binding is None:
            known = ", ".join(r.role.value for r in theme.roles)
            err(snode, f"в шаблоне нет макета для роли «{slide.layout}» (есть: {known})")
        for j, block in enumerate(slide.blocks, start=1):
            bnode = f"{snode}/block:{block.id or j}"
            parsed = parse(bnode, block.type, block.params, block.type_version)
            if parsed is None:
                continue
            plugin, params = parsed
            if plugin.target_kind != BlockTargetKind.SLOT:
                err(bnode, f"блок «{block.type}» заполняет фигуру слайда-образца, а этот слайд — из макета")
                continue
            needs_ok(bnode, plugin, params)
            slot = block.slot or DEFAULT_SLOT
            if binding is not None and binding.slot(slot) is None:
                known = ", ".join(s.name for s in binding.slots)
                err(bnode, f"у макета «{binding.layout_name}» нет области «{slot}» (есть: {known})")
    return issues


def markers_params(scenario: ScenarioSpec, slide: SlideSpec, example: TemplateSlideInfo | None) -> dict[str, Any]:
    """Параметры блока меток слайда-образца: привязки слайда и те привязки всей презентации,
    чьи метки есть на этом слайде (тогда слайд зависит только от своих показателей)."""
    names = set(example.marker_names) if example is not None else None
    common = {k: v for k, v in scenario.markers.items() if names is None or _marker_name(k) in names}
    return {"bindings": dict(slide.markers), "common": common}


def _validate_example(
    scenario: ScenarioSpec,
    slide: SlideSpec,
    snode: str,
    theme: ThemeManifest | None,
    err: Any,
    parse: Any,
    needs_ok: Any,
) -> None:
    assert slide.example is not None
    example = theme.slide(slide.example.id) if theme is not None else None
    if theme is not None and example is None:
        known = ", ".join(f"{s.slide_id} (слайд {s.number})" for s in theme.slides) or "нет"
        err(snode, f"в шаблоне нет слайда с id {slide.example.id} (есть: {known})")
        return
    filled: dict[int, str] = {}
    for j, block in enumerate(slide.blocks, start=1):
        bnode = f"{snode}/block:{block.id or j}"
        parsed = parse(bnode, block.type, block.params, block.type_version)
        if parsed is None:
            continue
        plugin, params = parsed
        if plugin.target_kind == BlockTargetKind.SLOT:
            err(
                bnode,
                f"блок «{block.type}» ставится на слайд из макета; на слайде-образце — метки, chart_fill, table_fill",
            )
            continue
        if plugin.target_kind == BlockTargetKind.MARKERS:
            err(bnode, "метки слайда-образца задаются полем markers слайда")
            continue
        if block.shape is None:
            err(bnode, f"у блока «{block.type}» не указана фигура: shape — id графика или таблицы на слайде")
            continue
        sid = block.shape.id
        if sid in filled:
            err(bnode, f"фигуру {sid} уже заполняет блок {filled[sid]}")
            continue
        filled[sid] = block.id or str(j)
        if example is not None:
            for p in plugin.check(params, example, sid):
                err(bnode, p, soft=True)
        needs_ok(bnode, plugin, params)

    parsed = parse(f"{snode}/markers", MARKERS_BLOCK, markers_params(scenario, slide, example))
    if parsed is not None:
        plugin, params = parsed
        if example is not None and (example.markers or slide.markers):
            for p in plugin.check(params, example, None):
                err(f"{snode}/markers", p, soft=True)
        needs_ok(f"{snode}/markers", plugin, params)

    if example is None:
        return
    shapes = [(BlockTargetKind.CHART, c.shape_id, c.shape_name) for c in example.charts]
    shapes += [(BlockTargetKind.TABLE, t.shape_id, t.shape_name) for t in example.tables]
    for kind, sid, name in shapes:
        if sid not in filled and sid not in slide.keep:
            err(
                f"{snode}/shape:{sid}",
                f"{_KIND_NAMES[kind]} «{name}» (id {sid}) не заполнен: привяжите к набору или добавьте id в keep, "
                "иначе останутся демо-данные шаблона",
                soft=True,
            )
    known_ids = {sid for _, sid, _ in shapes}
    for sid in slide.keep:
        if sid not in known_ids:
            err(f"{snode}/shape:{sid}", f"в keep id {sid}, а такого графика или таблицы на слайде нет")


# --- сборка -----------------------------------------------------------------------------


@dataclass
class _Ctx:
    period: Period
    scenario_name: str
    node: str
    preview: bool = False
    issues: list[Issue] = field(default_factory=list)
    kept: set[tuple[int, str]] = field(default_factory=set)

    def warn(self, message: str) -> None:
        self.issues.append(Issue(level=IssueLevel.WARNING, node=self.node, message=message))

    def keep_marker(self, shape_id: int, name: str) -> None:
        self.kept.add((shape_id, name))


@dataclass
class _Planned:
    number: int
    spec: SlideSpec
    slide: Any
    slide_id: int
    origin: int | None


def _layouts_by_key(prs: Any) -> dict[str, Any]:
    out = {}
    for master in prs.slide_masters:
        lst = master._element.find(f"{P_NS}sldLayoutIdLst")
        for entry, layout in zip(list(lst) if lst is not None else [], master.slide_layouts, strict=False):
            out[str(entry.get("id"))] = layout
    return out


def _error_marker(slide: Any, geometry: Any, message: str) -> None:
    box = slide.shapes.add_textbox(
        Emu(geometry.x), Emu(geometry.y), Emu(geometry.cx), Emu(min(geometry.cy, int(Pt(60))))
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = f"Ошибка: {message}"
    for r in tf.paragraphs[0].runs:
        r.font.color.rgb = ERROR_COLOR
        r.font.size = Pt(14)


def _remove_empty_placeholders(slide: Any) -> None:
    """Плейсхолдеры-надписи без своего текста (только с подсказкой, видной в режиме правки).
    Плейсхолдеры с графиком, таблицей или картинкой остаются."""
    for ph in list(slide.placeholders):
        el = ph._element
        if el.tag != f"{P}sp":
            continue
        if ph.has_text_frame and ph.text_frame.text.strip():
            continue
        el.getparent().remove(el)


def _drop_placeholders(slide: Any, indexes: list[int]) -> None:
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in indexes:
            ph._element.getparent().remove(ph._element)


def _all_shapes(shapes: Any) -> dict[int, Any]:
    """Фигуры python-pptx по id, заходя в группы."""
    out: dict[int, Any] = {}
    for sh in shapes:
        out[sh.shape_id] = sh
        if isinstance(sh, GroupShape):
            out.update(_all_shapes(sh.shapes))
    return out


def _slide_title(slide: Any, n: int) -> str:
    try:
        title = slide.shapes.title
    except (KeyError, AttributeError):
        title = None
    text = title.text_frame.text.strip() if title is not None and title.has_text_frame else ""
    return " ".join(text.split()) or f"Слайд {n}"


def _missing_data(needs: Any, data: BlockData, failed: dict[str, NodeStatus]) -> str | None:
    deps = [f"dataset:{d}" for d in sorted(needs.datasets)] + [f"metric:{m}" for m in sorted(needs.metrics)]
    broken = [d for d in deps if d in failed]
    if broken:
        root = failed[broken[0]].blocked_by or broken[0]
        return f"зависит от «{root}», в котором ошибка"
    # Итоги движка могли прийти не целиком (например, render запущен отдельно на сохранённых
    # файлах): блоку без своих данных — понятная ошибка, а не KeyError.
    absent = [f"набора «{d}»" for d in sorted(needs.datasets) if d not in data.datasets]
    absent += [f"показателя «{m}»" for m in sorted(needs.metrics) if m not in data.metrics]
    if absent:
        return f"в итогах расчёта нет {', '.join(absent)}"
    return None


def _fill_layout_slide(
    item: _Planned,
    binding: Any,
    scenario: ScenarioSpec,
    engine: EngineResult,
    registry: PluginRegistry,
    data: BlockData,
    failed: dict[str, NodeStatus],
    issues: list[Issue],
    preview: bool,
) -> list[str]:
    slide, snode = item.slide, f"slide:{item.number}"
    phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    errors: list[str] = []
    for j, block in enumerate(item.spec.blocks, start=1):
        bnode = f"{snode}/block:{block.id or j}"
        slot = binding.slot(block.slot or DEFAULT_SLOT)
        assert slot is not None
        plugin = registry.block(block.type)
        params = plugin.parse_params(block.params, block.type_version)
        ctx = _Ctx(period=engine.period, scenario_name=scenario.name, node=bnode, preview=preview)
        target = BlockTarget(
            slide=slide,
            geometry=slot.geometry,
            placeholder=phs.get(slot.placeholder_idx) if slot.placeholder_idx is not None else None,
            slot=slot.name,
        )
        try:
            missing = _missing_data(plugin.data_needs(params), data, failed)
            if missing:
                raise ValueError(missing)
            plugin.render(target, params, data, ctx)
        except Exception as e:
            msg = "; ".join(e.problems) if isinstance(e, BlockError) else _first_line(e)
            errors.append(f"блок {block.id or j} ({block.type}): {msg}")
            ctx.issues.append(Issue(level=IssueLevel.ERROR, node=bnode, message=msg))
            _error_marker(slide, slot.geometry, msg)
        issues.extend(ctx.issues)
    _remove_empty_placeholders(slide)
    return errors


def _fill_example_slide(
    item: _Planned,
    example: TemplateSlideInfo,
    scenario: ScenarioSpec,
    engine: EngineResult,
    registry: PluginRegistry,
    data: BlockData,
    failed: dict[str, NodeStatus],
    issues: list[Issue],
    kept: set[tuple[int, str]],
    preview: bool,
) -> list[str]:
    """Заполнить слайд-образец. Возвращает все места, которые не заполнились."""
    slide, snode = item.slide, f"slide:{item.number}"
    shapes = _all_shapes(slide.shapes)
    geometry = Geometry(x=0, y=0, cx=0, cy=0)
    problems: list[str] = []
    finish: dict[str, tuple[BlockPlugin, list[tuple[BlockTarget, Any]]]] = {}
    jobs: list[tuple[str, str, BlockPlugin, Any, BlockTarget]] = []

    if example.markers or item.spec.markers:
        plugin = registry.block(MARKERS_BLOCK)
        params = plugin.parse_params(markers_params(scenario, item.spec, example))
        jobs.append((f"{snode}/markers", "метки", plugin, params, BlockTarget(slide, geometry, example=example)))
    for j, block in enumerate(item.spec.blocks, start=1):
        plugin = registry.block(block.type)
        params = plugin.parse_params(block.params, block.type_version)
        assert block.shape is not None
        sid = block.shape.id
        where = f"{_KIND_NAMES.get(plugin.target_kind, 'фигура')} id {sid}"
        shape = shapes.get(sid)
        if shape is None:
            problems.append(f"{where}: такой фигуры на слайде нет")
            continue
        target = BlockTarget(slide, geometry, shape=shape, example=example)
        jobs.append((f"{snode}/block:{block.id or j}", where, plugin, params, target))

    for node, where, plugin, params, target in jobs:
        ctx = _Ctx(period=engine.period, scenario_name=scenario.name, node=node, preview=preview)
        try:
            missing = _missing_data(plugin.data_needs(params), data, failed)
            if missing:
                raise BlockError([missing])
            plugin.render(target, params, data, ctx)
            finish.setdefault(plugin.name, (plugin, []))[1].append((target, params))
        except BlockError as e:
            problems += [f"{where}: {p}" for p in e.problems]
        except Exception as e:
            problems.append(f"{where}: {_first_line(e)}")
        issues.extend(ctx.issues)
        kept |= ctx.kept

    for plugin, items in finish.values():
        ctx = _Ctx(period=engine.period, scenario_name=scenario.name, node=snode, preview=preview)
        try:
            plugin.finish_slide(slide, items, ctx)
        except BlockError as e:
            problems += e.problems
        except Exception as e:
            problems.append(_first_line(e))
        issues.extend(ctx.issues)
    _remove_empty_placeholders(slide)
    return problems


def _leftovers(slide: Any, kept: set[tuple[int, str]]) -> list[str]:
    """Где на слайде, в заметках и графиках осталось «{{», кроме меток «оставить метку»."""
    places: list[str] = []
    tree = slide._element.find(f"{P}cSld/{P}spTree")
    for sh in iter_shapes(tree):
        if sh.el.tag == f"{P}grpSp":
            continue
        text = all_text(sh.el)
        if "{{" not in text:
            continue
        names = {name for i, name in kept if i == sh.id}
        rest = "".join(
            piece if k % 2 == 0 else ("" if piece.strip() in names else "{{" + piece + "}}")
            for k, piece in enumerate(LOOSE_RE.split(text))
        )
        if "{{" in rest:
            places.append(f"«{sh.name}» (id {sh.id})")
    if slide.has_notes_slide and "{{" in all_text(slide.notes_slide._element):
        places.append("заметки к слайду")
    for sh in _all_shapes(slide.shapes).values():
        if getattr(sh, "has_chart", False) and sh.has_chart:
            cs = sh.chart._chartSpace
            texts = [t.text or "" for t in xp(cs, ".//a:t | .//c:v")]
            if any("{{" in t for t in texts):
                places.append(f"график «{sh.name}» (id {sh.shape_id})")
    return places


def build_presentation(
    scenario: ScenarioSpec,
    theme: ThemeManifest,
    engine: EngineResult,
    registry: PluginRegistry,
    output: str | Path,
    *,
    preview: bool = False,
    only: int | None = None,
) -> RenderResult:
    """Собрать .pptx по слайдам сценария, манифесту шаблона и итогам движка.

    ``preview`` — пробная сборка для картинки слайда: непривязанные и пустые метки остаются в
    тексте и подсвечиваются, незаполненные места — предупреждения. ``only`` — собрать только
    этот слайд сценария (номер среди включённых, с единицы)."""
    result = RenderResult()
    issues = validate_slides(scenario, registry, theme, preview=preview)
    if any(i.level == IssueLevel.ERROR for i in issues):
        result.issues = issues
        return result
    # Предупреждения проверки показывает тот, кто её вызвал (run, превью): здесь не повторяются.
    issues = []

    prs = Presentation(theme.pptx_path)
    template = slide_entries(prs)
    template_slides = dict(template)
    layouts = _layouts_by_key(prs)
    specs = list(enumerate((s for s in scenario.slides if s.enabled), start=1))
    if only is not None:
        if not 1 <= only <= len(specs):
            raise AgenError(ErrorCode.SPEC_REFERENCE, f"В сценарии нет слайда {only} (включённых: {len(specs)})")
        specs = [specs[only - 1]]

    # Сначала все новые слайды и копии — пока слайды шаблона нетронуты, потом удаление
    # лишних и порядок. Заполнение — после: копия берётся с чистого слайда-образца.
    plan: list[_Planned] = []
    used: set[int] = set()
    for n, spec in specs:
        if spec.example is not None:
            sid = spec.example.id
            src = template_slides[sid]
            if sid in used:
                slide = duplicate_slide(prs, src)
                new_id = int(prs.slides._sldIdLst[-1].get("id"))
            else:
                slide, new_id = src, sid
                used.add(sid)
            plan.append(_Planned(n, spec, slide, new_id, sid))
        else:
            binding = theme.role(spec.layout or "")
            assert binding is not None  # проверено в validate_slides
            layout = layouts.get(binding.layout_key)
            if layout is None:
                raise AgenError(ErrorCode.LAYOUT_MISSING, f"В рабочей копии шаблона нет макета {binding.layout_key}")
            slide = add_slide(prs, layout)
            _drop_placeholders(slide, binding.drop_placeholders)
            plan.append(_Planned(n, spec, slide, int(prs.slides._sldIdLst[-1].get("id")), None))
    remove_slides(prs, {sid for sid, _ in template} - used)
    reorder_slides(prs, [p.slide_id for p in plan])

    failed = {n.id: n for n in engine.nodes if n.state != NodeState.OK}
    data = BlockData(datasets=engine.datasets, metrics=engine.metrics)
    stop: list[Issue] = []
    kept: dict[int, set[tuple[int, str]]] = {}
    for item in plan:
        snode = f"slide:{item.number}"
        if item.spec.example is not None:
            example = theme.slide(item.spec.example.id)
            assert example is not None
            kept[item.slide_id] = set()
            problems = _fill_example_slide(
                item, example, scenario, engine, registry, data, failed, issues, kept[item.slide_id], preview
            )
            level = IssueLevel.WARNING if preview else IssueLevel.ERROR
            found = [Issue(level=level, node=snode, message=p) for p in problems]
            (issues if preview else stop).extend(found)
            errors = [] if preview else problems
        else:
            binding = theme.role(item.spec.layout or "")
            errors = _fill_layout_slide(item, binding, scenario, engine, registry, data, failed, issues, preview)
        result.nodes.append(
            NodeStatus(
                id=snode,
                kind=NodeKind.SLIDE,
                state=NodeState.ERROR if errors else NodeState.OK,
                message="; ".join(errors) or None,
            )
        )

    # Проверка после сборки: в слайдах, заметках и графиках не должно остаться «{{», кроме
    # меток, для которых выбрано «оставить метку». Такой отчёт не сохраняется. Слайды, где
    # блоки уже не заполнились, не проверяются: их места уже в списке.
    stopped = {i.node for i in stop}
    for item in plan:
        if f"slide:{item.number}" in stopped:
            continue
        places = _leftovers(item.slide, kept.get(item.slide_id, set()))
        if places:
            stop.append(
                Issue(
                    level=IssueLevel.ERROR,
                    node=f"slide:{item.number}",
                    message=f"осталась незаменённая метка «{{{{»: {', '.join(places)}",
                )
            )
    result.issues = issues + stop
    result.slides = len(prs.slides)
    if stop:
        return result

    origin = {p.slide_id: p.origin for p in plan}
    rewrite_sections(prs, origin)
    titles = [_slide_title(p.slide, i) for i, p in enumerate(plan, start=1)]
    rewrite_app_xml(prs, titles, len(template))
    drop_unused_comment_authors(prs)
    renumber_parts(prs)
    broken = check_package(prs)
    if broken:
        result.issues.append(
            Issue(level=IssueLevel.ERROR, message="файл отчёта собран с ошибками: " + "; ".join(broken[:10]))
        )
        return result

    props = prs.core_properties
    props.title = f"{scenario.name}: {engine.period.key}"
    props.author = "Autogenerator"
    props.last_modified_by = "Autogenerator"
    props.modified = datetime.now(UTC).replace(tzinfo=None)
    result.output_path = str(save_presentation(prs, Path(output), result.issues))
    return result


# --- сохранение ---------------------------------------------------------------------------

_RESERVED = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def sanitize_filename(name: str) -> str:
    """Имя файла, допустимое в Windows: без ``<>:"/\\|?*``, завершающих точек и пробелов и
    зарезервированных имён."""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).rstrip(" .")
    if not s:
        s = "Отчёт"
    if s.split(".")[0].upper() in _RESERVED:
        s = "_" + s
    return s


def _busy(path: Path) -> bool:
    if not path.exists():
        return False
    try:
        with path.open("ab"):
            return False
    except PermissionError:
        return True


def save_presentation(prs: Any, path: Path, issues: list[Issue]) -> Path:
    """Сохранить атомарно. Если файл открыт в PowerPoint, сохранить рядом как «… (2).pptx»."""
    path.parent.mkdir(parents=True, exist_ok=True)
    target = path
    n = 2
    while _busy(target):
        target = path.with_name(f"{path.stem} ({n}){path.suffix}")
        n += 1
        if n > 99:
            raise AgenError(ErrorCode.OUTPUT_BUSY, f"Файл {path.name} и его копии заняты")
    if target != path:
        issues.append(
            Issue(
                level=IssueLevel.WARNING,
                message=f"{path.name} открыт в другой программе; сохранено как {target.name}",
            )
        )
    # Временный файл рядом и переименование: полузаписанного отчёта не бывает. Файл создаёт
    # сам python-pptx, поэтому права — обычные, как у любого нового файла пользователя.
    tmp = target.with_name(f".{target.stem}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        prs.save(str(tmp))
        os.replace(tmp, target)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    return target
