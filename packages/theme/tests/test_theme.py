import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from autogenerator.contracts import AgenError, LayoutRole
from autogenerator.theme import import_template
from autogenerator.theme.importer import ZIP_RATIO_MIN_BYTES


@pytest.fixture
def template(tmp_path: Path) -> Path:
    """Шаблон python-pptx по умолчанию: стандартные макеты Office и один слайд с метками."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Продажи за {{Месяц}}"
    s.shapes.add_textbox(0, 0, 100, 100).text_frame.text = "Выручка {{ Выручка }} и снова {{Месяц}}"
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    return path


def test_roles_and_slots(template: Path, tmp_path: Path):
    m = import_template(template, tmp_path / "out")
    assert Path(m.pptx_path).exists()
    assert {r.role for r in m.roles} == set(LayoutRole)
    assert [s.name for s in m.role("title").slots] == ["title", "subtitle"]
    assert [s.name for s in m.role("title_and_content").slots] == ["title", "body"]
    left, right = (
        m.role("title_and_two_content").slot("left"),
        m.role("title_and_two_content").slot("right"),
    )
    assert left.geometry.x < right.geometry.x
    # У «только заголовка» область блоков — под заголовком.
    title_only = m.role("title_only")
    assert title_only.slot("body").geometry.y > title_only.slot("title").geometry.y
    assert title_only.slot("body").placeholder_idx is None
    # Финального макета в шаблоне по умолчанию нет: роль построена из титульного.
    assert m.role("final").derived and m.role("final").layout_key == m.role("title").layout_key
    assert m.notes == ["Роль «final» построена из макета «Title Slide»: финальный слайд на титульном макете"]


def test_template_slides_and_markers(template: Path, tmp_path: Path):
    m = import_template(template, tmp_path / "out")
    [slide] = m.slides
    assert slide.marker_names == ["Месяц", "Выручка"]
    box = slide.markers[1].shape_id
    assert [m.key for m in slide.markers][1:] == [f"Выручка@{box}#1", f"Месяц@{box}#1"]
    assert slide.number == 1 and slide.slide_id >= 256


def _retype(src: Path, dst: Path, content_type: str, extra: dict[str, bytes] | None = None) -> Path:
    main = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(main.encode(), content_type.encode())
            zout.writestr(item, data)
        for name, data in (extra or {}).items():
            zout.writestr(name, data)
    return dst


def test_potx_is_opened_as_presentation(template: Path, tmp_path: Path):
    potx = _retype(
        template,
        tmp_path / "t.potx",
        "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
    )
    m = import_template(potx, tmp_path / "out")
    assert any(".potx" in n for n in m.notes)
    Presentation(m.pptx_path)


def test_macros_are_rejected(template: Path, tmp_path: Path):
    pptm = _retype(
        template,
        tmp_path / "t.pptm",
        "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    )
    with pytest.raises(AgenError, match="макросы"):
        import_template(pptm, tmp_path / "out")
    sneaky = _retype(
        template,
        tmp_path / "t2.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        {"ppt/vbaProject.bin": b"x"},
    )
    with pytest.raises(AgenError, match="макросы"):
        import_template(sneaky, tmp_path / "out2")


def test_not_a_pptx(tmp_path: Path):
    f = tmp_path / "x.pptx"
    f.write_text("не презентация", encoding="utf-8")
    with pytest.raises(AgenError, match="не файл PowerPoint"):
        import_template(f, tmp_path / "out")
    with pytest.raises(AgenError, match="не найден"):
        import_template(tmp_path / "nope.pptx", tmp_path / "out")


def test_zip_bomb_is_rejected(template: Path, tmp_path: Path):
    bomb = tmp_path / "bomb.pptx"
    with zipfile.ZipFile(template) as zin, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("ppt/media/zeros.bin", b"\0" * (ZIP_RATIO_MIN_BYTES + 1))
    with pytest.raises(AgenError, match="zip-бомб"):
        import_template(bomb, tmp_path / "out")
