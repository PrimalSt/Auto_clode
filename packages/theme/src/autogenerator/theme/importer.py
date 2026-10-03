"""Импорт .pptx-шаблона: рабочая копия, макеты всех мастеров, роли макетов, слайды-образцы,
шрифты и проверка шаблона (ARCHITECTURE.md, раздел 6.5)."""

from __future__ import annotations

import hashlib
import re
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from autogenerator.contracts import AgenError, ErrorCode, TemplateSlideInfo, ThemeManifest

from .decorations import find_decorations
from .fonts import template_fonts
from .layouts import guess_roles, read_layouts
from .lint import lint
from .slides import slide_info

CT_PRESENTATION = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
CT_TEMPLATE = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
CT_SLIDESHOW = "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml"

ZIP_RATIO_LIMIT = 100
"""Часть файла, сжатая сильнее 100:1, — признак zip-бомбы (как у .xlsx, раздел 12)."""
ZIP_RATIO_MIN_BYTES = 10 << 20
ZIP_TOTAL_LIMIT = 2 << 30


def check_zip(path: Path) -> None:
    """Отклонить zip-бомбу: слишком сильно сжатую часть или слишком большой распакованный объём."""
    total = 0
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            total += info.file_size
            if (
                info.file_size >= ZIP_RATIO_MIN_BYTES
                and info.compress_size
                and info.file_size / info.compress_size > ZIP_RATIO_LIMIT
            ):
                raise AgenError(
                    ErrorCode.THEME_INVALID,
                    f"Шаблон {path.name} отклонён: часть {info.filename} сжата сильнее {ZIP_RATIO_LIMIT}:1 "
                    "(так выглядят zip-бомбы)",
                )
    if total > ZIP_TOTAL_LIMIT:
        raise AgenError(
            ErrorCode.THEME_INVALID,
            f"Шаблон {path.name} отклонён: в распакованном виде больше {ZIP_TOTAL_LIMIT >> 30} ГБ",
        )


def _main_content_type(path: Path) -> str:
    try:
        check_zip(path)
        with zipfile.ZipFile(path) as z:
            ct = z.read("[Content_Types].xml").decode("utf-8")
            names = set(z.namelist())
    except (zipfile.BadZipFile, KeyError) as e:
        raise AgenError(ErrorCode.THEME_INVALID, f"{path.name} — не файл PowerPoint") from e
    m = re.search(r'PartName="/ppt/presentation\.xml"\s+ContentType="([^"]+)"', ct) or re.search(
        r'ContentType="([^"]+)"\s+PartName="/ppt/presentation\.xml"', ct
    )
    if not m:
        raise AgenError(ErrorCode.THEME_INVALID, f"В {path.name} нет основной части презентации")
    if "macroEnabled" in m.group(1) or "ppt/vbaProject.bin" in names:
        raise AgenError(
            ErrorCode.THEME_INVALID,
            f"{path.name} содержит макросы; такие шаблоны не принимаются",
            hint="Сохраните шаблон в PowerPoint как «Презентация PowerPoint (.pptx)» без макросов.",
        )
    return m.group(1)


def make_working_copy(src: Path, dst: Path) -> list[str]:
    """Скопировать шаблон так, чтобы его открыл python-pptx (у .potx меняется тип основной части)."""
    notes: list[str] = []
    ct = _main_content_type(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    if ct == CT_PRESENTATION:
        shutil.copyfile(src, dst)
        return notes
    if ct not in (CT_TEMPLATE, CT_SLIDESHOW):
        raise AgenError(ErrorCode.THEME_INVALID, f"Неизвестный тип файла {src.name}: {ct}")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(ct.encode(), CT_PRESENTATION.encode())
            zout.writestr(item, data)
    notes.append("Шаблон .potx/.ppsx открыт как презентация: в рабочей копии заменён тип основной части")
    return notes


def import_template(
    path: str | Path,
    out_dir: str | Path,
    *,
    check: bool = True,
    roles: dict[str, str] | None = None,
    strict_roles: bool = False,
) -> ThemeManifest:
    """Импортировать шаблон: рабочая копия в ``out_dir/template.pptx`` и манифест с отчётом
    проверки (``check=False`` — без проверки, для быстрого просмотра). ``roles`` —
    подтверждённые роли макетов (роль → ключ макета); ``strict_roles`` — неподходящий макет
    подтверждения — ошибка, а не замечание."""
    src = Path(path)
    if not src.exists():
        raise AgenError(ErrorCode.FILE_NOT_FOUND, f"Шаблон не найден: {src}")
    work = Path(out_dir) / "template.pptx"
    notes = make_working_copy(src, work)
    try:
        prs = Presentation(str(work))
    except Exception as e:
        raise AgenError(ErrorCode.THEME_INVALID, f"python-pptx не открыл шаблон {src.name}: {e}") from e
    width, height = int(prs.slide_width or Emu(0)), int(prs.slide_height or Emu(0))
    pairs = read_layouts(prs)
    layouts = [info for info, _ in pairs]
    roles_, role_notes = guess_roles(layouts, width, height, roles, strict_roles)
    notes += role_notes
    keys = {layout.part.partname: info.key for info, layout in pairs}
    ids = [int(s.get("id")) for s in prs.slides._sldIdLst]
    slides: list[TemplateSlideInfo] = [
        slide_info(slide, number, sid, keys.get(slide.slide_layout.part.partname, ""))
        for number, (sid, slide) in enumerate(zip(ids, prs.slides, strict=True), start=1)
    ]
    roles_ = find_decorations(
        prs,
        roles_,
        {info.key: layout for info, layout in pairs},
        [(s.slide_id, slide, s.layout_key) for s, slide in zip(slides, prs.slides, strict=True)],
    )
    fonts = template_fonts(prs)
    report = lint(prs, slides, layouts, fonts, width, height) if check else []
    return ThemeManifest(
        source_path=str(src),
        pptx_path=str(work),
        sha256=hashlib.sha256(src.read_bytes()).hexdigest(),
        slide_width=width,
        slide_height=height,
        layouts=layouts,
        roles=roles_,
        slides=slides,
        fonts=fonts,
        lint=report,
        notes=notes,
    )
