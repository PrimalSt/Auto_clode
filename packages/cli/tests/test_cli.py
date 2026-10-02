from pathlib import Path

from pptx import Presentation
from typer.testing import CliRunner

from autogenerator.cli.main import app
from autogenerator.contracts import PreviewResult

ROOT = Path(__file__).resolve().parents[3]
SCENARIO = ROOT / "examples" / "sales" / "scenario.yaml"
runner = CliRunner()


def test_run_example(tmp_path: Path):
    out = tmp_path / "отчёт.pptx"
    r = runner.invoke(app, ["run", str(SCENARIO), "-o", str(out)])
    assert r.exit_code == 0, r.output
    assert "Отчётный период: 2026-03" in r.output
    assert "Готово:" in r.output
    assert len(Presentation(str(out)).slides) == 11


def test_run_with_explicit_inputs_and_period(tmp_path: Path):
    data = ROOT / "examples" / "sales" / "data"
    r = runner.invoke(
        app,
        [
            "run",
            str(SCENARIO),
            "--output-dir",
            str(tmp_path),
            "--period",
            "2026-01",
            "-i",
            f"sales={data / 'sales' / 'Продажи_2026-01.csv'}",
            "-i",
            f"plan={data / 'plan' / 'План_2026-Q1.xlsx'}",
        ],
    )
    # Январь — первый месяц: прошлого месяца в истории нет, это предупреждение, а не ошибка.
    assert r.exit_code == 0, r.output
    assert (tmp_path / "Отчёт_продажи_2026-01.pptx").exists()
    assert "нет загрузок входа «sales» за" in r.output


def test_validate(tmp_path: Path):
    r = runner.invoke(app, ["validate", str(SCENARIO)])
    assert r.exit_code == 0 and "Сценарий в порядке" in r.output
    bad = tmp_path / "scenario.yaml"
    bad.write_text(SCENARIO.read_text(encoding="utf-8").replace("fn: sum", "fn: summ", 1), encoding="utf-8")
    (tmp_path / "sources.yaml").write_text(
        (SCENARIO.parent / "sources.yaml").read_text(encoding="utf-8"), encoding="utf-8"
    )
    r = runner.invoke(
        app,
        ["validate", str(bad), "--theme", str(ROOT / "examples" / "templates" / "synthetic.pptx")],
    )
    assert r.exit_code == 1
    assert "summ" in r.output


def test_inspect():
    f = ROOT / "examples" / "sales" / "data" / "sales" / "Продажи_2026-02.csv"
    r = runner.invoke(app, ["inspect", str(f)])
    assert r.exit_code == 0, r.output
    assert "кодировка cp1251, разделитель ';'" in r.output
    assert "Сумма, руб." in r.output


def test_modules():
    r = runner.invoke(app, ["modules"])
    assert r.exit_code == 0, r.output
    assert "ОШИБКА" not in r.output
    for name in ("csv", "xlsx", "dedupe", "quarter_to_date", "sum", "chart"):
        assert name in r.output


def test_preview_step_dataset_metric_and_sample():
    r = runner.invoke(app, ["preview", str(SCENARIO), "sales/positive_only"])
    assert r.exit_code == 0, r.output
    assert "Превью input:sales/step:positive_only · период 2026-03" in r.output
    assert "positive_only        (filter)" in r.output and "← превью после этого шага" in r.output
    assert "net_of_vat" not in r.output  # шаги после выбранного не выполняются
    r = runner.invoke(app, ["preview", str(SCENARIO), "plan_fact", "--json"])
    assert r.exit_code == 0, r.output
    res = PreviewResult.model_validate_json(r.output)
    assert [c.name for c in res.columns] == ["region", "fact", "plan", "done"]
    assert res.total_rows == 5 and not res.approximate
    r = runner.invoke(app, ["preview", str(SCENARIO), "metric:revenue", "--period", "2026-02"])
    assert r.exit_code == 0, r.output
    assert "период 2026-02" in r.output and "revenue_prev_change_pct = " in r.output
    r = runner.invoke(app, ["preview", str(SCENARIO), "sales", "--sample", "3", "--rows", "2"])
    assert r.exit_code == 0, r.output
    assert "выборка ≈1/3 (sales: order_no)" in r.output and "Строк: ≈" in r.output


def test_preview_unknown_node():
    r = runner.invoke(app, ["preview", str(SCENARIO), "dataset:nope"])
    assert r.exit_code == 1
    assert "Нет узла «dataset:nope»" in r.output


def test_bad_input_argument():
    r = runner.invoke(app, ["run", str(SCENARIO), "-i", "sales.csv"])
    assert r.exit_code == 2


TEMPLATE = ROOT / "examples" / "templates" / "synthetic.pptx"


def test_theme_check_and_scaffold(tmp_path: Path):
    r = runner.invoke(app, ["theme", "check", str(TEMPLATE)])
    # В синтетическом шаблоне есть слайд «Черновик» с ошибками проверки — код выхода 1.
    assert r.exit_code == 1, r.output
    assert "Роли макетов:" in r.output and "[ошибка] слайд 7" in r.output
    out = tmp_path / "slides.yaml"
    r = runner.invoke(app, ["theme", "scaffold", str(TEMPLATE), "-o", str(out)])
    assert r.exit_code == 0, r.output
    text = out.read_text(encoding="utf-8")
    assert "Месяц: period.month" in text and "type: chart_fill" in text and "type: table_fill" in text


def test_preview_slide(tmp_path: Path):
    out = tmp_path / "slide.pptx"
    r = runner.invoke(app, ["preview", str(SCENARIO), "slide:2", "-o", str(out)])
    assert r.exit_code == 0, r.output
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    assert any("Итоги, март 2026" in sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    r = runner.invoke(app, ["preview", str(SCENARIO), "slide:99", "-o", str(out)])
    assert r.exit_code != 0
