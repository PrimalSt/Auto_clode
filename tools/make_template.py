"""Синтетический шаблон оформления для примеров и тестов: ``examples/templates/synthetic.pptx``.

Корпоративный шаблон пользователя в репозиторий не кладётся. Этот шаблон повторяет его
особенности (ARCHITECTURE.md, раздел 13), но с выдуманным содержанием про продажи:

- два мастера; во втором — макет «Последний» для финального слайда, копия макета «Только
  заголовок» без ``preserve`` (на ней стоят слайды-образцы, а такая же копия с ``preserve`` есть
  в первом мастере) и макет «Заголовок и объект», у плейсхолдеров которого нет геометрии ни в
  макете, ни в мастере;
- слайды-образцы со свободно расставленными надписями поверх макета: метки ``{{…}}``, разрезанные
  PowerPoint на прогоны с разными языками и флагом ошибки орфографии; имена меток с ``+``,
  ``-`` и ``=``; одно имя на разных слайдах (``{{Доля}}`` на слайдах 2 и 3 значит разное);
  после меток в шаблоне уже написаны единицы («%», «млн ₽»);
- комбинированный график (столбцы с накоплением и линия на второй оси со скрытой осью),
  пара наложенных друг на друга графиков, круговая диаграмма с настройками секторов,
  100%-гистограмма с надписями напротив каждой полосы, пустая таблица со стилем;
- служебные объекты и теги think-cell (условные), заметки к слайду, список разделов, авторы
  примечаний без примечаний;
- слайд «Черновик» с ошибками, которые находит проверка шаблона: вписанный вручную год, ссылка
  на номер слайда, метка, разрезанная переносом строки, метка внутри графика, повторяющиеся
  имена фигур, фигура за краем слайда. В примерах сценариев он не используется.

Графики заполнены демо-данными («Категория 1», «Ряд 1»), как в настоящем шаблоне.
"""

from __future__ import annotations

import copy
import io
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, XmlPart
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlideLayoutPart, SlideMasterPart
from pptx.presentation import Presentation as PresentationT
from pptx.util import Emu, Inches, Pt

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)
ACCENTS = {
    "accent1": "1F4E79",
    "accent2": "F28C28",
    "accent3": "2E8B57",
    "accent4": "7F7F7F",
    "accent5": "5B9BD5",
    "accent6": "C00000",
}
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}
A = f"{{{NS['a']}}}"
P = f"{{{NS['p']}}}"
R = f"{{{NS['r']}}}"
C = f"{{{NS['c']}}}"
CT_TAGS = "application/vnd.openxmlformats-officedocument.presentationml.tags+xml"
CT_AUTHORS = "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"
RT_AUTHORS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors"
INK = RGBColor(0x2B, 0x2F, 0x47)
GREY = RGBColor(0x7F, 0x7F, 0x7F)
FONT = "Arial"


# --- текст: прогоны с языками, как их режет PowerPoint ----------------------------------------


def _run(p: Any, text: str, lang: str, size: float | None = None, bold: bool | None = None, err: bool = False) -> Any:
    r = p.add_run()
    r.text = text
    rpr = r._r.get_or_add_rPr()
    rpr.set("lang", lang)
    if err:
        rpr.set("err", "1")
    if size:
        r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold
    r.font.color.rgb = INK
    r.font.name = FONT
    return r


def marker(name: str, pieces: int = 3) -> list[tuple[str, str]]:
    """Метка, разрезанная на прогоны: «{{» и «}}» — en-US, имя — ru-RU; при ``pieces > 3`` имя
    ещё делится, как бывает после правки текста в PowerPoint."""
    parts = [("{{", "en-US")]
    if pieces <= 3 or len(name) < 2:
        parts.append((name, "ru-RU"))
    else:
        cut = max(1, len(name) // 2)
        parts += [(name[:cut], "en-US"), (name[cut:], "ru-RU")]
    parts.append(("}}", "en-US"))
    if pieces >= 5:
        parts = [("{", "en-US"), ("{", "ru-RU"), *parts[1:]]
    return parts


def text_box(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    paragraphs: list[list[tuple[str, str]]],
    size: float = 18,
    bold: bool = False,
    align: Any = PP_ALIGN.LEFT,
    wrap: bool = True,
    name: str | None = None,
) -> Any:
    """Надпись из абзацев; абзац — список прогонов (текст, язык)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT if not wrap else MSO_AUTO_SIZE.NONE
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, lang in runs:
            cyr = any("а" <= ch.lower() <= "я" for ch in text)
            _run(p, text, lang, size, bold, err=cyr and lang == "en-US")
    return box


# --- графики ---------------------------------------------------------------------------------


def _demo(categories: int, series: int, seed: int = 0) -> CategoryChartData:
    cd = CategoryChartData(number_format="General")
    cd.categories = [f"Категория {i + 1}" for i in range(categories)]
    for s in range(series):
        cd.add_series(f"Ряд {s + 1}", [round(2 + ((i + s + seed) * 1.7) % 4, 1) for i in range(categories)])
    return cd


def _labels(plot: Any, fmt: str, position: Any | None = None, size: float = 11) -> None:
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.show_value = True
    dl.font.size = Pt(size)
    if position is not None:
        dl.position = position


def _el(tag: str, **attrs: str) -> etree._Element:
    prefix, local = tag.split(":")
    # Префиксы p, a, r, c объявлены в корне частей; остальные (p14) — на самом элементе,
    # иначе lxml назовёт их ns0.
    nsmap = None if prefix in ("p", "a", "r", "c") else {prefix: NS[prefix]}
    e = etree.Element(f"{{{NS[prefix]}}}{local}", nsmap=nsmap)
    for k, v in attrs.items():
        e.set(k, v)
    return e


def _sub(parent: etree._Element, tag: str, **attrs: str) -> etree._Element:
    e = _el(tag, **attrs)
    parent.append(e)
    return e


def _make_combo(chart: Any, line_format: str = "0%", hidden_axis: bool = True) -> None:
    """Перенести последнюю серию столбцов в линию на второй оси значений (как в шаблоне
    пользователя: python-pptx комбинированные графики не строит)."""
    cs = chart._chartSpace
    plot_area = cs.find(f"{C}chart/{C}plotArea")
    bar = plot_area.find(f"{C}barChart")
    ser = bar.findall(f"{C}ser")[-1]
    bar.remove(ser)
    inv = ser.find(f"{C}invertIfNegative")
    if inv is not None:
        ser.remove(inv)
    sppr = ser.find(f"{C}spPr")
    if sppr is not None:
        ser.remove(sppr)
    sppr = _el("c:spPr")
    ln = _sub(sppr, "a:ln", w="28575", cap="rnd")
    _sub(_sub(ln, "a:solidFill"), "a:schemeClr", val="accent2")
    anchor = ser.find(f"{C}tx")
    if anchor is None:
        anchor = ser.find(f"{C}order")
    anchor.addnext(sppr)
    mk = _el("c:marker")
    _sub(mk, "c:symbol", val="circle")
    _sub(mk, "c:size", val="7")
    sppr.addnext(mk)
    ser.find(f"{C}val").addnext(_el("c:smooth", val="0"))

    line = _el("c:lineChart")
    _sub(line, "c:grouping", val="standard")
    _sub(line, "c:varyColors", val="0")
    line.append(ser)
    dls = _sub(line, "c:dLbls")
    _sub(dls, "c:numFmt", formatCode=line_format, sourceLinked="0")
    _sub(dls, "c:spPr").append(_el("a:noFill"))
    _sub(dls, "c:dLblPos", val="t")
    for flag, v in (
        ("showLegendKey", "0"),
        ("showVal", "1"),
        ("showCatName", "0"),
        ("showSerName", "0"),
        ("showPercent", "0"),
        ("showBubbleSize", "0"),
    ):
        _sub(dls, f"c:{flag}", val=v)
    _sub(line, "c:marker", val="1")
    _sub(line, "c:axId", val="50010")
    _sub(line, "c:axId", val="50011")
    bar.addnext(line)

    axes = [e for e in plot_area if etree.QName(e).localname.endswith("Ax")]
    last = axes[-1]
    cat = _el("c:catAx")
    _sub(cat, "c:axId", val="50010")
    _sub(_sub(cat, "c:scaling"), "c:orientation", val="minMax")
    _sub(cat, "c:delete", val="1")
    _sub(cat, "c:axPos", val="b")
    _sub(cat, "c:majorTickMark", val="none")
    _sub(cat, "c:minorTickMark", val="none")
    _sub(cat, "c:tickLblPos", val="nextTo")
    _sub(cat, "c:crossAx", val="50011")
    _sub(cat, "c:crosses", val="autoZero")
    _sub(cat, "c:auto", val="1")
    _sub(cat, "c:lblAlgn", val="ctr")
    _sub(cat, "c:lblOffset", val="100")
    _sub(cat, "c:noMultiLvlLbl", val="0")
    val = _el("c:valAx")
    _sub(val, "c:axId", val="50011")
    _sub(_sub(val, "c:scaling"), "c:orientation", val="minMax")
    _sub(val, "c:delete", val="0")
    _sub(val, "c:axPos", val="r")
    _sub(val, "c:numFmt", formatCode="General", sourceLinked="1")
    _sub(val, "c:majorTickMark", val="none")
    _sub(val, "c:minorTickMark", val="none")
    _sub(val, "c:tickLblPos", val="nextTo")
    if hidden_axis:
        # Ось линии «спрятана» белым текстом — так сделано в шаблоне пользователя.
        txpr = _sub(val, "c:txPr")
        _sub(txpr, "a:bodyPr")
        _sub(txpr, "a:lstStyle")
        para = _sub(txpr, "a:p")
        defrpr = _sub(_sub(para, "a:pPr"), "a:defRPr", sz="900")
        _sub(_sub(defrpr, "a:solidFill"), "a:srgbClr", val="FFFFFF")
        _sub(para, "a:endParaRPr", lang="ru-RU")
    _sub(val, "c:crossAx", val="50010")
    _sub(val, "c:crosses", val="max")
    _sub(val, "c:crossBetween", val="between")
    last.addnext(cat)
    cat.addnext(val)


def add_chart(slide: Any, kind: Any, x: float, y: float, w: float, h: float, cd: CategoryChartData) -> Any:
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    chart = frame.chart
    chart.font.size = Pt(11)
    chart.font.name = FONT
    return frame


def _pie_points(chart: Any, n: int) -> None:
    """Секторы с собственными цветами и подписями «категория; доля» (dPt и dLbl на каждый
    сектор, как в шаблоне пользователя)."""
    plot = chart.plots[0]
    plot.vary_by_categories = True
    colors = list(ACCENTS.values())
    ser = plot.series[0]
    for i in range(n):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(colors[i % len(colors)])
        pt.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    for dl in ser._element.findall(f"{C}dLbls/{C}dLbl"):
        for flag, v in (
            ("showLegendKey", "0"),
            ("showVal", "0"),
            ("showCatName", "1"),
            ("showSerName", "0"),
            ("showPercent", "1"),
            ("showBubbleSize", "0"),
        ):
            _sub(dl, f"c:{flag}", val=v)
        _sub(dl, "c:separator").text = "; "


# --- служебные части: теги, авторы примечаний, разделы -----------------------------------------


def add_tags(slide_part: Any, shape: Any, tags: dict[str, str]) -> None:
    """Тег PowerPoint на фигуре (``p:custDataLst`` → часть ``/ppt/tags/tagN.xml``): так помечает
    свои фигуры think-cell, а приложение — выгруженные слайды (F-422)."""
    package = slide_part.package
    root = _el("p:tagLst")
    for k, v in tags.items():
        _sub(root, "p:tag", name=k, val=v)
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    part = XmlPart.load(package.next_partname("/ppt/tags/tag%d.xml"), CT_TAGS, package, blob)
    rid = slide_part.relate_to(part, RT.TAGS)
    nvpr = shape._element.find(f".//{P}nvPr")
    cust = _el("p:custDataLst")
    tag = _sub(cust, "p:tags")
    tag.set(f"{R}id", rid)
    nvpr.append(cust)


def add_comment_authors(prs: PresentationT) -> None:
    """Список авторов примечаний без единого примечания: в отчёт он не должен попасть."""
    root = _el("p:cmAuthorLst")
    for i, (name, ini) in enumerate((("Автор шаблона", "АШ"), ("Дизайнер", "Д"))):
        _sub(root, "p:cmAuthor", id=str(i), name=name, initials=ini, lastIdx="1", clrIdx=str(i))
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    part = Part(PackURI("/ppt/commentAuthors.xml"), CT_AUTHORS, prs.part.package, blob)
    prs.part.relate_to(part, RT_AUTHORS)


def add_sections(prs: PresentationT, sections: list[tuple[str, list[int]]]) -> None:
    """Список разделов (``p14:sectionLst``): python-pptx его не ведёт, сборка переписывает его сама."""
    ids = [int(s.get("id")) for s in prs.slides._sldIdLst]
    pres = prs.part._element
    ext_lst = pres.find(f"{P}extLst")
    if ext_lst is None:
        ext_lst = _sub(pres, "p:extLst")
    ext = _sub(ext_lst, "p:ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    lst = _sub(ext, "p14:sectionLst")
    for name, numbers in sections:
        sec = _sub(lst, "p14:section", name=name, id=f"{{{str(uuid.uuid5(uuid.NAMESPACE_URL, name)).upper()}}}")
        sl = _sub(sec, "p14:sldIdLst")
        for n in numbers:
            _sub(sl, "p14:sldId", id=str(ids[n - 1]))


# --- оформление: цвета темы, полоса на мастере, второй мастер ----------------------------------


def _scale_xfrm(root: etree._Element, fx: float) -> None:
    """Растянуть по горизонтали все фигуры макета: шаблон python-pptx рассчитан на 4:3."""
    for off in root.iter(f"{A}off"):
        off.set("x", str(int(int(off.get("x", "0")) * fx)))
    for ext in root.iter(f"{A}ext"):
        if ext.get("cx") is not None:
            ext.set("cx", str(int(int(ext.get("cx", "0")) * fx)))


def _theme_part(master_part: Any) -> Any:
    return next(r.target_part for r in master_part.rels.values() if r.reltype == RT.THEME)


def _theme_colors(prs: PresentationT) -> None:
    theme_part = _theme_part(prs.slide_master.part)
    root = etree.fromstring(theme_part.blob)
    for name, rgb in ACCENTS.items():
        el = root.find(f".//{A}clrScheme/{A}{name}")
        if el is not None:
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{A}srgbClr", val=rgb)
    root.find(f".//{A}clrScheme").set("name", "Autogenerator")
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _master_band(prs: PresentationT) -> None:
    """Полоса акцентного цвета сверху и подпись в подвале на мастере: так видно, что отчёт
    собран в шаблоне, а не в пустой презентации."""
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    band = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDTH, Inches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(ACCENTS["accent1"])
    band.line.fill.background()
    band.name = "Полоса шаблона"
    note = tmp.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(6), Inches(0.35))
    note.name = "Подпись шаблона"
    note.text_frame.text = "Синтетический шаблон Autogenerator"
    note.text_frame.word_wrap = True
    note.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    run = note.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(10)
    run.font.color.rgb = GREY
    tree = prs.slide_master.shapes._spTree
    for sh in (band, note):
        tree.append(copy.deepcopy(sh._element))
    _drop_last_slide(prs)


def _drop_last_slide(prs: PresentationT) -> None:
    sld = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


def _logo(shapes: Any, x: float, y: float) -> None:
    """Условный логотип: группа из квадрата и подписи. Такие же группы есть на макете
    «Последний» второго мастера."""
    group = shapes.add_group_shape()
    group.name = "Логотип"
    sq = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.5), Inches(0.5))
    sq.fill.solid()
    sq.fill.fore_color.rgb = RGBColor.from_string(ACCENTS["accent2"])
    sq.line.fill.background()
    word = group.shapes.add_textbox(Inches(x + 0.6), Inches(y + 0.05), Inches(2.2), Inches(0.4))
    word.text_frame.text = "Синтетика"
    r = word.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENTS["accent1"])


def _second_master(prs: PresentationT) -> dict[str, Any]:
    """Второй мастер (в русском PowerPoint — «образец слайдов») со своими макетами. Возвращает
    макеты по именам. python-pptx мастера не добавляет, поэтому части собираются вручную."""
    package = prs.part.package
    m1 = prs.slide_master
    m1_part = m1.part
    theme1 = _theme_part(m1_part)
    theme2 = Part(package.next_partname("/ppt/theme/theme%d.xml"), theme1.content_type, package, theme1.blob)

    root = copy.deepcopy(m1._element)
    lst = root.find(f"{P}sldLayoutIdLst")
    for e in list(lst):
        lst.remove(e)
    # На втором мастере только заголовок: у плейсхолдеров «объект» макетов этого мастера
    # геометрии нет нигде (как у макета «Заголовок и объект» в шаблоне пользователя).
    tree = root.find(f"{P}cSld/{P}spTree")
    for sp in list(tree):
        ph = sp.find(f".//{P}nvPr/{P}ph")
        if ph is not None and ph.get("type") not in ("title",):
            tree.remove(sp)
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    m2_part = SlideMasterPart.load(
        package.next_partname("/ppt/slideMasters/slideMaster%d.xml"), m1_part.content_type, package, blob
    )
    m2_part.relate_to(theme2, RT.THEME)
    lst = m2_part._element.find(f"{P}sldLayoutIdLst")
    pres = prs.part._element
    # id мастеров и макетов общие на весь файл: повтор PowerPoint считает повреждением.
    ids = [int(e.get("id")) for e in pres.iter(f"{P}sldMasterId")]
    ids += [int(e.get("id")) for e in m1_part._element.iter(f"{P}sldLayoutId")]
    next_id = max(ids) + 1
    m_entry = _el("p:sldMasterId", id=str(next_id))
    m_entry.set(f"{R}id", prs.part.relate_to(m2_part, RT.SLIDE_MASTER))
    pres.find(f"{P}sldMasterIdLst").append(m_entry)

    by_name = {lay.name: lay for lay in m1.slide_layouts}
    out: dict[str, Any] = {}

    def add_layout(src_name: str, name: str, preserve: bool, edit: Any = None) -> None:
        nonlocal next_id
        el = copy.deepcopy(by_name[src_name]._element)
        el.find(f"{P}cSld").set("name", name)
        if preserve:
            el.set("preserve", "1")
        elif "preserve" in el.attrib:
            del el.attrib["preserve"]
        if edit is not None:
            edit(el)
        part = SlideLayoutPart.load(
            package.next_partname("/ppt/slideLayouts/slideLayout%d.xml"),
            by_name[src_name].part.content_type,
            package,
            etree.tostring(el, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        part.relate_to(m2_part, RT.SLIDE_MASTER)
        next_id += 1
        entry = _el("p:sldLayoutId", id=str(next_id))
        entry.set(f"{R}id", m2_part.relate_to(part, RT.SLIDE_LAYOUT))
        lst.append(entry)
        out[name] = part.slide_layout

    def final(el: etree._Element) -> None:
        tree = el.find(f"{P}cSld/{P}spTree")
        for sp in list(tree):
            if sp.find(f".//{P}nvPr/{P}ph") is not None:
                tree.remove(sp)

    def no_geometry(el: etree._Element) -> None:
        tree = el.find(f"{P}cSld/{P}spTree")
        for sp in list(tree):
            ph = sp.find(f".//{P}nvPr/{P}ph")
            if ph is None:
                continue
            if ph.get("type") in ("dt", "ftr", "sldNum"):
                tree.remove(sp)
                continue
            xfrm = sp.find(f"{P}spPr/{A}xfrm")
            if xfrm is not None:
                xfrm.getparent().remove(xfrm)

    add_layout("Title Only", "Только заголовок", preserve=False)
    add_layout("Blank", "Последний", preserve=True, edit=final)
    add_layout("Title and Content", "Заголовок и объект", preserve=False, edit=no_geometry)
    # Логотип на макете «Последний»: тот же, что на обложке. У макета нет add_group_shape,
    # поэтому группа собирается на временном слайде и копируется.
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    _logo(tmp.shapes, 5.2, 3.2)
    out["Последний"].shapes._spTree.append(copy.deepcopy(tmp.shapes[-1]._element))
    _drop_last_slide(prs)
    return out


# --- слайды-образцы ----------------------------------------------------------------------------


def _title(slide: Any, parts: list[tuple[str, str]], name: str = "Заголовок 1") -> Any:
    return text_box(slide, 0.5, 0.35, 12.3, 0.8, [parts], size=28, bold=True, align=PP_ALIGN.CENTER, name=name)


def _cover(prs: PresentationT) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Продажи"
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    text_box(
        s,
        3.5,
        4.4,
        6.3,
        0.8,
        [[*marker("Месяц", 4), (" ", "ru-RU"), *marker("Год")]],
        size=28,
        align=PP_ALIGN.CENTER,
        name="Период",
    )
    _logo(s.shapes, 5.2, 6.2)


def _summary(prs: PresentationT, layout: Any) -> None:
    """Слайд 2: итоги месяца с комбинированным графиком (столбцы с накоплением и линия на
    второй оси)."""
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _title(s, [("Итоги", "ru-RU"), (", ", "en-US"), *marker("Месяц", 5), (" ", "ru-RU"), *marker("Год")])
    text_box(
        s,
        0.6,
        1.4,
        3.6,
        1.1,
        [[*marker("Выручка"), (" млн ₽", "ru-RU")], [("выручка без НДС", "ru-RU")]],
        size=24,
        bold=True,
        name="Заголовок 2",
    )
    text_box(
        s,
        4.6,
        1.4,
        3.6,
        1.1,
        [[*marker("Прирост+"), ("%", "ru-RU")], [("к прошлому месяцу", "ru-RU")]],
        size=24,
        bold=True,
        name="Заголовок 2",
    )
    text_box(
        s, 8.6, 1.4, 4.0, 1.1, [[*marker("Доля"), ("% плана выполнено", "ru-RU")]], size=24, bold=True, name="План"
    )
    frame = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, 0.6, 2.7, 12.1, 4.2, _demo(3, 3))
    chart = frame.chart
    frame.name = "Диаграмма 3"
    _labels(chart.plots[0], "#,##0.0")
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    _make_combo(chart, "0%")
    s.notes_slide.notes_text_frame.text = "Заметка шаблона: проверить выручку перед отправкой."
    add_tags(s.part, s.shapes[1], {"TCLAYOUT": "Синтетика"})  # условный тег think-cell


def _regions(prs: PresentationT, layout: Any) -> None:
    """Слайд 3: круговая диаграмма с настройками секторов; {{Доля}} здесь — доля Москвы."""
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _title(s, [("Регионы", "ru-RU"), (", ", "en-US"), *marker("Месяц"), (" ", "ru-RU"), *marker("Год")])
    text_box(
        s,
        0.6,
        1.6,
        4.0,
        1.4,
        [[*marker("Доля"), ("%", "ru-RU")], [("выручки — у первого региона", "ru-RU")]],
        size=24,
        bold=True,
        name="Доля",
    )
    frame = add_chart(s, XL_CHART_TYPE.PIE, 5.0, 1.3, 7.6, 5.6, _demo(6, 1))
    frame.name = "Диаграмма 5"
    chart = frame.chart
    chart.has_legend = False
    _pie_points(chart, 6)


def _plan(prs: PresentationT, layout: Any) -> None:
    """Слайд 4: 100%-гистограмма и надписи с метками напротив каждой полосы."""
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _title(s, [("План по регионам", "ru-RU"), (", ", "en-US"), *marker("Месяц"), (" ", "ru-RU"), *marker("Год")])
    cd = _demo(5, 2)
    frame = add_chart(s, XL_CHART_TYPE.BAR_STACKED_100, 0.6, 1.3, 9.6, 5.8, cd)
    frame.name = "Диаграмма 7"
    chart = frame.chart
    _labels(chart.plots[0], "0%", XL_LABEL_POSITION.CENTER)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    # Надписи напротив полос: подогнаны под пять категорий шаблона.
    for i in range(5):
        y = 1.65 + i * 0.93
        text_box(s, 10.4, y, 2.4, 0.5, [[*marker(f"р{i + 1}="), (" п.п.", "ru-RU")]], size=16, name=f"Надпись {i + 1}")


def _pair(prs: PresentationT, layout: Any) -> None:
    """Слайд 5: два графика, наложенные друг на друга и выровненные по категориям, как пара
    на слайдах оплат в шаблоне пользователя."""
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _title(s, [("Заказы", "ru-RU"), (", ", "en-US"), *marker("Месяц"), (" ", "ru-RU"), *marker("Год")])
    text_box(s, 0.6, 1.3, 5.0, 0.6, [[*marker("+заказы"), (" заказов", "ru-RU")]], size=20, bold=True, name="Заказы")
    text_box(
        s, 6.6, 1.3, 5.0, 0.6, [[*marker("возвраты-"), (" возвратов", "ru-RU")]], size=20, bold=True, name="Возвраты"
    )
    bottom = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, 0.6, 2.0, 12.1, 4.9, _demo(3, 2, seed=1))
    bottom.name = "Диаграмма 9"
    _labels(bottom.chart.plots[0], "#,##0")
    bottom.chart.has_legend = False
    top = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, 0.6, 2.0, 12.1, 4.9, _demo(3, 3, seed=2))
    top.name = "Диаграмма 10"
    _labels(top.chart.plots[0], "#,##0")
    top.chart.has_legend = False
    _make_combo(top.chart, "#,##0")
    # Верхний график прозрачный: сквозь него виден нижний.
    cs = top.chart._chartSpace
    sppr = _el("c:spPr")
    sppr.append(_el("a:noFill"))
    cs.find(f"{C}chart").addnext(sppr)
    plot_area = cs.find(f"{C}chart/{C}plotArea")
    psppr = _el("c:spPr")
    psppr.append(_el("a:noFill"))
    plot_area.append(psppr)


def _table(prs: PresentationT, layout: Any) -> None:
    """Слайд 6: пустая таблица со стилем на макете без геометрии и служебный объект."""
    s = prs.slides.add_slide(layout)
    title = s.shapes.title
    title.left, title.top, title.width, title.height = Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8)
    title.text = "Выручка регионов по месяцам"
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    frame = s.shapes.add_table(5, 4, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.4))
    frame.name = "Таблица 4"
    tbl = frame.table
    for j, h in enumerate(["Регион", "Месяц 1", "Месяц 2", "Месяц 3"]):
        tbl.cell(0, j).text = h
    tbl.columns[0].width = Inches(3.7)
    for j in range(1, 4):
        tbl.columns[j].width = Inches(2.8)
    text_box(s, 0.6, 6.6, 8.0, 0.4, [[("Источник: CRM, выручка без НДС, млн ₽", "ru-RU")]], size=12, name="Сноска")
    # Пустой служебный объект, как у think-cell на слайдах с таблицей и графиком когорт.
    ole = s.shapes.add_ole_object(
        io.BytesIO(b"\x00" * 64), "TCLayout.ActiveDocument.1", Inches(0.1), Inches(0.1), Inches(0.1), Inches(0.1)
    )
    ole.name = "think-cell data - do not delete"


def _draft(prs: PresentationT, layout: Any) -> None:
    """Слайд 7: ошибки, которые находит проверка шаблона."""
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _title(s, [("Черновик", "ru-RU")])
    _title(s, [("Повтор имени", "ru-RU")])  # вторая фигура «Заголовок 1»
    text_box(
        s,
        0.6,
        1.4,
        8,
        0.6,
        [[("Рост к ", "ru-RU"), *marker("Месяц"), (" 2025", "ru-RU")]],
        size=18,
        name="Вписанный год",
    )
    text_box(s, 0.6, 2.1, 8, 0.6, [[("Подробнее — на слайде 2", "ru-RU")]], size=18, name="Ссылка на слайд")
    box = text_box(s, 0.6, 2.8, 8, 0.9, [[("{{Ме", "ru-RU")]], size=18, name="Метка с переносом")
    p = box.text_frame.paragraphs[0]
    p._p.append(_el("a:br"))
    _run(p, "сяц}}", "ru-RU", 18)
    text_box(s, 12.2, 4.0, 2.0, 0.6, [marker("Год")], size=18, name="За краем")
    frame = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.6, 4.0, 6.0, 3.0, _demo(3, 1))
    frame.name = "Диаграмма 12"
    frame.chart.has_title = True
    frame.chart.chart_title.text_frame.text = "Выручка {{Год}}"


def _final(prs: PresentationT, layout: Any) -> None:
    s = prs.slides.add_slide(layout)
    text_box(
        s, 3.5, 4.3, 6.3, 0.8, [[("Спасибо!", "ru-RU")]], size=32, bold=True, align=PP_ALIGN.CENTER, name="Спасибо"
    )


def write_template(path: Path) -> str:
    prs = Presentation()
    fx = WIDTH / prs.slide_width
    prs.slide_width, prs.slide_height = Emu(WIDTH), Emu(HEIGHT)
    _scale_xfrm(prs.slide_master._element, fx)
    for layout in prs.slide_master.slide_layouts:
        _scale_xfrm(layout._element, fx)
    # Копия «Только заголовок» с preserve="1": её приложение берёт для новых слайдов.
    title_only = next(lay for lay in prs.slide_layouts if lay.name == "Title Only")
    title_only._element.set("preserve", "1")
    _theme_colors(prs)
    _master_band(prs)
    second = _second_master(prs)

    _cover(prs)
    _summary(prs, second["Только заголовок"])
    _regions(prs, second["Только заголовок"])
    _plan(prs, second["Только заголовок"])
    _pair(prs, second["Только заголовок"])
    _table(prs, second["Заголовок и объект"])
    _draft(prs, second["Только заголовок"])
    _final(prs, second["Последний"])
    add_sections(prs, [("Титул", [1]), ("Продажи", [2, 3, 4, 5, 6]), ("Служебное", [7, 8])])
    add_comment_authors(prs)

    props = prs.core_properties
    props.title = "Синтетический шаблон Autogenerator"
    props.author = "Autogenerator"
    props.created = props.modified = datetime(2026, 9, 30)
    props.last_modified_by = "Autogenerator"
    props.revision = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    layouts = sum(len(m.slide_layouts) for m in prs.slide_masters)
    return f"{len(prs.slide_masters)} мастера, {layouts} макетов, {len(prs.slides)} слайдов шаблона"


if __name__ == "__main__":
    import sys

    default = Path(__file__).resolve().parents[1] / "examples" / "templates" / "synthetic.pptx"
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    print(f"{out.name}: {write_template(out)}")
